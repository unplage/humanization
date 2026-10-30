#!/usr/bin/env python3
"""Generate Word reports for all antibody sequences in a FASTA file."""

import argparse
import os
import re
import sys

# Add repo root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
# Add current directory to path for local imports
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from evaluate_humanness import (
    classify_humanization_level,
    evaluate_sequence,
    evaluate_sequence_imgt,
    load_sequences,
)
from scripts.humanize.anarci_adapter import number_with_anarci_imgt
from scripts.humanize.imgt_numbering import IMGT_REGIONS, load_imgt_germline


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading_elm = cell._element.get_or_add_tcPr()
    shading = shading_elm.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color
    })
    shading_elm.append(shading)


def get_usan_class(fr_identity):
    """Get academic humanization estimate based on FR identity."""
    return classify_humanization_level(fr_identity)


def get_imgt_region(pos_label):
    """Get IMGT region from position label like 'H105' or 'H111A'."""
    num_str = ""
    for c in pos_label[1:]:
        if c.isdigit():
            num_str += c
        else:
            break

    if not num_str:
        return "Unknown"

    num = int(num_str)
    for region, (start, end) in IMGT_REGIONS.items():
        if start <= num <= end:
            return region
    return "Unknown"


def _cdrs_from_numbered(numbered):
    """Collect CDR sequences from a pipeline NumberedChain in sequence order."""
    return {
        region: "".join(r.aa for r in numbered.residues if r.region == region)
        for region in ("CDR1", "CDR2", "CDR3")
    }


def get_cdr_sequences(sequence, chain_type, numbered=None):
    """Get CDR sequences using ANARCI IMGT numbering (pipeline regions)."""
    if numbered is None:
        numbered = number_with_anarci_imgt(sequence, chain_type)
    if not numbered:
        return None
    return _cdrs_from_numbered(numbered)


def get_kabat_cdr_sequences(sequence, chain_type, numbered=None):
    """Get CDR sequences using Kabat numbering (pipeline regions)."""
    from scripts.humanize.numbering import number_heavy, number_light

    if numbered is None:
        numbered = number_heavy(sequence) if chain_type == 'H' else number_light(sequence)
    return _cdrs_from_numbered(numbered)


def _region_stats(best_scores, key):
    return {
        'identity': best_scores.get(f'{key}_identity', 0),
        'match': best_scores.get(f'{key}_match', 0),
        'count': best_scores.get(f'n_{key}', 0),
    }


def evaluate_antibody(name, sequence, chain_type, db_dir="", numbered_kabat=None,
                      imgt_genes=None):
    """Evaluate a single antibody and return comprehensive results."""
    results = {}

    scored_kabat, numbered_kabat = evaluate_sequence(
        chain_type, sequence, db_dir=db_dir, numbered=numbered_kabat)
    if not scored_kabat:
        raise ValueError(f"no Kabat-scored germlines for {name!r}")
    best_gene_kabat, best_scores_kabat = scored_kabat[0]

    q_map = numbered_kabat.posmap()
    g_map = best_gene_kabat.numbered.posmap() if best_gene_kabat.numbered else {}
    diffs = []
    for pos in q_map:
        if pos in g_map and q_map[pos] != g_map[pos]:
            region = numbered_kabat.region_of(pos) or "?"
            diffs.append((pos, q_map[pos], g_map[pos], region))

    results['kabat'] = {
        'best_gene': best_gene_kabat.gene_id,
        'fr_identity': best_scores_kabat['fr_identity'],
        'cdr_identity': best_scores_kabat['cdr_identity'],
        'all_identity': best_scores_kabat['all_identity'],
        'n_diff': len(diffs),
        'n_query': best_scores_kabat.get('n_query', 0),
        'n_compared': best_scores_kabat.get('n_compared', 0),
        'coverage': best_scores_kabat.get('coverage', 0.0),
        'n_query_only': best_scores_kabat.get('n_query_only', 0),
        'n_germline_only': best_scores_kabat.get('n_germline_only', 0),
        'regions': {
            'FR1': _region_stats(best_scores_kabat, 'fr1'),
            'FR2': _region_stats(best_scores_kabat, 'fr2'),
            'FR3': _region_stats(best_scores_kabat, 'fr3'),
            'CDR1': _region_stats(best_scores_kabat, 'cdr1'),
            'CDR2': _region_stats(best_scores_kabat, 'cdr2'),
        },
        'diffs': diffs,
    }

    # IMGT evaluation (mandatory)
    scored_imgt, numbered_imgt = evaluate_sequence_imgt(
        chain_type, sequence, db_dir=db_dir, imgt_genes=imgt_genes)
    if not scored_imgt or not numbered_imgt:
        raise RuntimeError(f"IMGT evaluation failed for {name!r} - ANARCI or IMGT germline DB unavailable")

    best_gene_imgt, best_scores_imgt = scored_imgt[0]
    imgt_diffs = []
    q_map_imgt = numbered_imgt.posmap()
    g_map_imgt = getattr(best_gene_imgt, '_imgt_posmap', {}) or {}
    for pos in q_map_imgt:
        if pos in g_map_imgt and q_map_imgt[pos] != g_map_imgt[pos]:
            region = get_imgt_region(pos)
            imgt_diffs.append((pos, q_map_imgt[pos], g_map_imgt[pos], region))

    results['imgt'] = {
        'available': True,
        'best_gene': best_gene_imgt.gene_id,
        'fr_identity': best_scores_imgt['fr_identity'],
        'cdr_identity': best_scores_imgt['cdr_identity'],
        'all_identity': best_scores_imgt['all_identity'],
        'n_diff': len(imgt_diffs),
        'usan_class': get_usan_class(best_scores_imgt['fr_identity']),
        'n_query': best_scores_imgt.get('n_query', 0),
        'n_compared': best_scores_imgt.get('n_compared', 0),
        'coverage': best_scores_imgt.get('coverage', 0.0),
        'n_query_only': best_scores_imgt.get('n_query_only', 0),
        'n_germline_only': best_scores_imgt.get('n_germline_only', 0),
        'regions': best_scores_imgt.get('imgt_region_stats', {}),
        'diffs': imgt_diffs,
    }
    results['cdrs'] = _cdrs_from_numbered(numbered_imgt)

    results['kabat_cdrs'] = _cdrs_from_numbered(numbered_kabat)
    return results


def add_antibody_section(doc, name, sequence, chain_type, results):
    """Add detailed section for one antibody."""
    doc.add_heading(f'{name} ({chain_type} chain)', level=2)

    doc.add_paragraph(f'Chain type: {"Heavy" if chain_type == "H" else "Light"}')
    doc.add_paragraph(f'Length: {len(sequence)} amino acids')
    doc.add_paragraph(f'Sequence: {sequence}')

    kabat = results['kabat']
    doc.add_heading('Kabat Numbering', level=3)
    doc.add_paragraph(f'Best germline: {kabat["best_gene"]}')
    doc.add_paragraph(f'FR identity: {kabat["fr_identity"]*100:.1f}%')
    doc.add_paragraph(f'CDR identity: {kabat["cdr_identity"]*100:.1f}%')
    doc.add_paragraph(f'Overall identity: {kabat["all_identity"]*100:.1f}%')
    doc.add_paragraph(f'Positions differing: {kabat["n_diff"]}')
    doc.add_paragraph(
        f'Alignment coverage: {kabat["n_compared"]}/{kabat["n_query"]} '
        f'({kabat["coverage"]*100:.1f}%); query-only {kabat["n_query_only"]}, '
        f'germline-only {kabat["n_germline_only"]}'
    )

    doc.add_heading('Kabat Region Breakdown', level=4)
    for region, stats in kabat['regions'].items():
        doc.add_paragraph(
            f'{region}: {stats.get("identity", 0)*100:.1f}% '
            f'({stats.get("match", 0)}/{stats.get("count", 0)})',
            style='List Bullet'
        )

    if results.get('kabat_cdrs'):
        doc.add_heading('Kabat CDR Sequences', level=4)
        for region, seq in results['kabat_cdrs'].items():
            doc.add_paragraph(f'{region}: {seq}', style='List Bullet')

    if kabat['diffs']:
        doc.add_heading('Kabat Differing Positions', level=4)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'

        headers = ['Position', 'Query', 'Germline', 'Region']
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

        for pos, q_aa, g_aa, region in kabat['diffs'][:30]:
            row = table.add_row()
            row.cells[0].text = pos
            row.cells[1].text = q_aa
            row.cells[2].text = g_aa
            row.cells[3].text = region

    imgt = results['imgt']
    doc.add_heading('IMGT Numbering (WHO/INN/USAN)', level=3)
    doc.add_paragraph(f'Best germline: {imgt["best_gene"]}')
    doc.add_paragraph(f'FR identity: {imgt["fr_identity"]*100:.1f}%')
    doc.add_paragraph(f'CDR identity: {imgt["cdr_identity"]*100:.1f}%')
    doc.add_paragraph(f'Overall identity: {imgt["all_identity"]*100:.1f}%')
    doc.add_paragraph(f'Humanization estimate: {imgt["usan_class"]}')
    doc.add_paragraph(
        f'Alignment coverage: {imgt["n_compared"]}/{imgt["n_query"]} '
        f'({imgt["coverage"]*100:.1f}%); query-only {imgt["n_query_only"]}, '
        f'germline-only {imgt["n_germline_only"]}'
    )
    doc.add_paragraph(
        '(Note: Official USAN/INN naming requires knowledge of manufacturing process)',
        style='Intense Quote'
    )

    if imgt['regions']:
        doc.add_heading('IMGT Region Breakdown', level=4)
        for region, stats in imgt['regions'].items():
            if isinstance(stats, dict):
                doc.add_paragraph(
                    f'{region}: {stats.get("identity", 0)*100:.1f}% '
                    f'({stats.get("match", 0)}/{stats.get("count", 0)})',
                    style='List Bullet'
                )
            else:
                doc.add_paragraph(f'{region}: {stats}', style='List Bullet')

    if results.get('cdrs'):
        doc.add_heading('IMGT CDR Sequences', level=4)
        for region, seq in results['cdrs'].items():
            doc.add_paragraph(f'{region}: {seq}', style='List Bullet')

    if imgt['diffs']:
        doc.add_heading('IMGT Differing Positions', level=4)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'

        headers = ['Position', 'Query', 'Germline', 'Region']
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True

        for pos, q_aa, g_aa, region in imgt['diffs'][:30]:
            row = table.add_row()
            row.cells[0].text = pos
            row.cells[1].text = q_aa
            row.cells[2].text = g_aa
            row.cells[3].text = region

    doc.add_page_break()


def create_combined_report(all_results, output_path):
    """Create a combined Word report with summary table and all antibody sections."""
    doc = Document()

    title = doc.add_heading('Antibody Panel - Humanization Evaluation Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading('Overview', level=1)
    doc.add_paragraph(f'Total antibodies evaluated: {len(all_results)}')
    doc.add_paragraph('Evaluation methods: Kabat numbering + IMGT numbering (WHO/INN/USAN)')
    doc.add_paragraph('Excluded regions: CDR3 (donor) and FR4 (human J gene)')
    doc.add_paragraph('IMGT evaluation is mandatory (requires ANARCI + IMGT germline DB)')
    doc.add_paragraph('Note: Official USAN/INN naming requires manufacturing process knowledge')

    doc.add_page_break()

    doc.add_heading('Summary', level=1)

    doc.add_heading('Kabat Numbering Summary', level=2)
    table_kabat = doc.add_table(rows=1, cols=6)
    table_kabat.style = 'Table Grid'
    table_kabat.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_kabat = ['Antibody', 'Chain', 'Best Germline', 'FR id', 'CDR id', '# Diff']
    for i, header in enumerate(headers_kabat):
        cell = table_kabat.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, '4472C4')
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for name, sequence, chain_type, results in all_results:
        kabat = results['kabat']
        row = table_kabat.add_row()
        row.cells[0].text = name
        row.cells[1].text = 'VH' if chain_type == 'H' else 'VL'
        row.cells[2].text = kabat['best_gene']
        row.cells[3].text = f"{kabat['fr_identity']*100:.1f}%"
        row.cells[4].text = f"{kabat['cdr_identity']*100:.1f}%"
        row.cells[5].text = str(kabat['n_diff'])

    doc.add_paragraph()

    doc.add_heading('IMGT Numbering Summary (WHO/INN/USAN)', level=2)
    table_imgt = doc.add_table(rows=1, cols=7)
    table_imgt.style = 'Table Grid'
    table_imgt.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_imgt = ['Antibody', 'Chain', 'Best Germline', 'FR id', 'CDR id', '# Diff', 'Estimate']
    for i, header in enumerate(headers_imgt):
        cell = table_imgt.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        set_cell_shading(cell, '548235')
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    for name, sequence, chain_type, results in all_results:
        imgt = results['imgt']
        row = table_imgt.add_row()
        row.cells[0].text = name
        row.cells[1].text = 'VH' if chain_type == 'H' else 'VL'
        row.cells[2].text = imgt['best_gene']
        row.cells[3].text = f"{imgt['fr_identity']*100:.1f}%"
        row.cells[4].text = f"{imgt['cdr_identity']*100:.1f}%"
        row.cells[5].text = str(imgt['n_diff'])
        row.cells[6].text = imgt.get('usan_class', 'N/A')

    doc.add_paragraph()

    for name, sequence, chain_type, results in all_results:
        add_antibody_section(doc, name, sequence, chain_type, results)

    doc.save(output_path)
    print(f"  总报告已保存: {output_path}")


def _safe_report_stem(name, used):
    stem = re.sub(r'[^A-Za-z0-9._-]+', '_', name).strip('._-') or 'antibody'
    candidate = stem
    suffix = 2
    while candidate.lower() in used:
        candidate = f'{stem}_{suffix}'
        suffix += 1
    used.add(candidate.lower())
    return candidate


def main():
    parser = argparse.ArgumentParser(description='Generate Word humanization reports')
    parser.add_argument('fasta', nargs='?', help='FASTA file with sequences')
    parser.add_argument('--input', '-i', dest='input_path',
                        help='FASTA file with sequences (alternative to positional)')
    parser.add_argument('--outdir', default='outputs/reports',
                        help='Directory for generated .docx files')
    parser.add_argument('--db-dir', help='Germline database directory')
    args = parser.parse_args()

    fasta_path = args.input_path or args.fasta or 'data/examples/antibody_panel.fasta'
    if not os.path.exists(fasta_path):
        print(f"Error: FASTA file not found: {fasta_path}", file=sys.stderr)
        return 1

    try:
        sequences, errors = load_sequences(fasta_path, skip_invalid=True)
    except (OSError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    for name, message in errors:
        print(f"  [ERROR] {name}: {message}", file=sys.stderr)
    if not sequences:
        print("Error: no valid antibody sequences to evaluate", file=sys.stderr)
        return 1

    os.makedirs(args.outdir, exist_ok=True)

    # Check ANARCI availability (required for IMGT)
    try:
        from scripts.humanize.anarci_adapter import number_with_anarci_imgt
        # Test ANARCI with a dummy sequence
        test_result = number_with_anarci_imgt("EVQLVESGGGLVQPGGSLRLSCAASGFTFTDYTMHWVRQAPGKGLEWVARIYPTNGYTRYADSVKGRFTISADTSKNTAYLQMNSLRAEDTAVYYCARMDYWGQGTLVTVSS", "H")
        if test_result is None:
            print("ERROR: ANARCI IMGT numbering failed. Please install ANARCI for IMGT evaluation:", file=sys.stderr)
            print("  pip install anarci", file=sys.stderr)
            return 1
    except ImportError:
        print("ERROR: ANARCI not installed. Required for IMGT numbering (WHO/INN/USAN standard).", file=sys.stderr)
        print("  Install with: pip install anarci", file=sys.stderr)
        return 1
    except Exception as e:
        print(f"ERROR: ANARCI IMGT numbering failed: {e}", file=sys.stderr)
        return 1

    # Load IMGT germline database
    imgt_genes = load_imgt_germline(args.db_dir or "")
    if not imgt_genes:
        print(
            "ERROR: IMGT germline database unavailable. Required for IMGT evaluation.",
            file=sys.stderr,
        )
        print("  Expected: data/germline/abnumber_human_imgt.json", file=sys.stderr)
        return 1

    print(f"正在评估 {len(sequences)} 个抗体序列...")
    print("  [INFO] IMGT evaluation enabled (ANARCI + IMGT germline DB)")

    all_results = []
    used_names = set()

    for name, chain_type, sequence, numbered_kabat in sequences:
        print(f"  评估 {name}...")
        try:
            results = evaluate_antibody(
                name, sequence, chain_type,
                db_dir=args.db_dir or "",
                numbered_kabat=numbered_kabat,
                imgt_genes=imgt_genes,
            )
        except Exception as exc:
            message = str(exc) or exc.__class__.__name__
            errors.append((name, message))
            print(f"  [ERROR] {name}: {message}", file=sys.stderr)
            continue

        all_results.append((name, sequence, chain_type, results))

        stem = _safe_report_stem(name, used_names)
        output_path = os.path.join(args.outdir, f'{stem}_report.docx')
        print(f"  生成 {name} 报告...")
        try:
            doc = Document()
            title = doc.add_heading(f'{name} - Humanization Evaluation Report', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_heading('Overview', level=1)
            doc.add_paragraph(f'Antibody: {name}')
            doc.add_paragraph(f'Chain type: {"Heavy" if chain_type == "H" else "Light"}')
            doc.add_paragraph(f'Sequence: {sequence}')
            doc.add_paragraph(f'Length: {len(sequence)} amino acids')

            doc.add_heading('Evaluation Criteria', level=2)
            if chain_type == 'H':
                doc.add_paragraph('Kabat numbering (VH): FR1 (1-30) + CDR1 (31-35) + FR2 (36-49) + CDR2 (50-65) + FR3 (66-94) + CDR3 (95-102) + FR4 (103-113)')
            else:
                doc.add_paragraph('Kabat numbering (VL): FR1 (1-23) + CDR1 (24-34) + FR2 (35-49) + CDR2 (50-56) + FR3 (57-88) + CDR3 (89-97) + FR4 (98-107)')
            doc.add_paragraph('IMGT numbering (same for VH/VL): FR1 (1-26) + CDR1 (27-38) + FR2 (39-55) + CDR2 (56-65) + FR3 (66-104) + CDR3 (105-117) + FR4 (118-128)')
            doc.add_paragraph('Excluded: CDR3 (from donor antibody) and FR4 (from human J gene)')
            doc.add_paragraph('Humanization estimate (Academic convention): Based on sequence identity, not official USAN/INN naming')
            doc.add_paragraph('Note: Official USAN/INN naming is based on the technology used (transgenic mouse, CDR grafting, phage display, etc.)')

            add_antibody_section(doc, name, sequence, chain_type, results)
            doc.save(output_path)
            print(f"  报告已保存: {output_path}")
        except Exception as exc:
            message = str(exc) or exc.__class__.__name__
            errors.append((name, message))
            print(f"  [ERROR] report {name}: {message}", file=sys.stderr)

    if not all_results:
        print("Error: no antibody reports were generated", file=sys.stderr)
        return 1

    print("\n  生成总报告...")
    combined_path = os.path.join(args.outdir, 'combined_report.docx')
    try:
        create_combined_report(all_results, combined_path)
    except Exception as exc:
        message = str(exc) or exc.__class__.__name__
        errors.append(('combined_report', message))
        print(f"  [ERROR] combined report: {message}", file=sys.stderr)

    if errors:
        print(f"\n报告生成完成，但有 {len(errors)} 个失败记录。")
        return 1

    print("\n所有报告生成完成！")
    return 0


if __name__ == "__main__":
    sys.exit(main())
