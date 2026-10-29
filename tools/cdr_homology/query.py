#!/usr/bin/env python3
"""CDR Homology Analysis -- compare antibody CDRs against a clinical database.

Usage:
    # Single query (VH + VL)
    python3 tools/cdr_homology/query.py \
        --vh QVQLVQSGAEVKKPGASVKVSCKAS... --vl DIQMTQSPSSLSASVGDRVTITC... \
        --top 5 --output outputs/cdr_homology/

    # From FASTA (VH/VL paired by header suffix)
    python3 tools/cdr_homology/query.py \
        --input query.fasta --top 5 --output outputs/cdr_homology/

Exit codes:
    0  success
    1  fatal error (database missing, no queries, unreadable input)
    2  partial success: some query chains could not be numbered (reports are
       still written, but the affected CDRs are marked N/A and listed under
       "Errors" in the report)
"""

from __future__ import annotations

import argparse
import csv
import os
import statistics
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import cdr_core as core  # noqa: E402

# Backwards-compatible names: the numbering/CDR helpers moved into cdr_core.
_extract_cdrs_from_posmap = core.extract_cdrs_from_posmap
extract_cdrs = core.extract_cdrs
_number_anarci = core.number_sequence_imgt


# ---------------------------------------------------------------------------
# Query input parsing
# ---------------------------------------------------------------------------

_HEAVY_MARKERS = ("VH", "HEAVY", "IGH")
_LIGHT_MARKERS = ("VL", "LIGHT", "KAPPA", "LAMBDA", "IGK", "IGL")


def _header_chain_hint(header: str):
    """Return ``'H'``/``'L'`` when the header states a chain type, else None."""
    token = header.split()[0].upper()
    if "|" in token:
        token = token.split("|")[1]
    # Check the light markers first: "VL" contains no "VH", but a header such
    # as "VL_Heavy" would otherwise match the heavy test first.
    for marker in _LIGHT_MARKERS:
        if marker in token:
            return "L"
    for marker in _HEAVY_MARKERS:
        if marker in token:
            return "H"
    return None


def _classify_record(header: str, sequence: str):
    """Return ``('H'|'L', 'header'|'anarci'|'default')`` for a FASTA record."""
    hint = _header_chain_hint(header)
    if hint:
        return hint, "header"
    detected = core.detect_chain_type(sequence)
    if detected:
        return detected, "anarci"
    # Last resort: a heavy chain is longer on average.  Flagged as a guess.
    return ("H" if len(sequence) > 120 else "L"), "default"


def parse_query_inputs(args) -> list:
    """Return ``[(name, vh_seq, vl_seq, meta), ...]`` from CLI args."""
    if args.vh or args.vl:
        name = core.base_name(args.name or "query_1")
        return [(name, args.vh, args.vl, {"name_source": "--name"})]

    records = core.parse_fasta(args.input)
    heavy, light = {}, {}
    for header, sequence in records:
        if not sequence:
            continue
        chain_type, source = _classify_record(header, sequence)
        slot = heavy if chain_type == "H" else light
        slot[header] = (sequence, source)

    order = list(dict.fromkeys([h for h, _ in records]))
    queries, used = [], set()
    for header in order:
        if header in used:
            continue
        base = core.base_name(header)
        vh = heavy.get(header, (None, None))
        vl = light.get(header, (None, None))
        vh_seq, vh_src = vh
        vl_seq, vl_src = vl
        if vh_seq and not vl_seq:
            for other, (seq, src) in light.items():
                if other not in used and core.base_name(other) == base:
                    vl_seq, vl_src = seq, src
                    used.add(other)
                    break
        elif vl_seq and not vh_seq:
            for other, (seq, src) in heavy.items():
                if other not in used and core.base_name(other) == base:
                    vh_seq, vh_src = seq, src
                    used.add(other)
                    break
        used.add(header)
        queries.append((base, vh_seq, vl_seq, {
            "vh_header": header if vh_seq else None,
            "vl_header": None,
            "chain_source": {"VH": vh_src, "VL": vl_src},
        }))
    return queries


# ---------------------------------------------------------------------------
# CDR homology search
# ---------------------------------------------------------------------------

def _iter_chain_cdrs(database: list, chain_key: str, cdr_name: str):
    """Yield ``(entry, subject_cdr, chain_source)`` for primary + bispecific."""
    for entry in database:
        chain_data = entry.get("chains", {}).get(chain_key)
        if chain_data:
            subject_cdr = chain_data.get(cdr_name, "")
            if subject_cdr:
                yield entry, subject_cdr, "primary"
        bispec_data = entry.get("bispecific", {}).get(chain_key)
        if bispec_data:
            subject_cdr = bispec_data.get(cdr_name, "")
            if subject_cdr:
                yield entry, subject_cdr, "bispecific"


def _make_hit(entry, subject_cdr, source, aln) -> dict:
    return {
        "therapeutic": entry.get("name", ""),
        "chain_source": source,
        "identity": aln["identity"],
        "identity_global": aln["identity_global"],
        "identity_over_query": aln["identity_over_query"],
        "coverage": aln["coverage"],
        "score": aln["score"],
        "score_norm": aln["score_norm"],
        "matches": aln["matches"],
        "aligned_len": aln["aligned_len"],
        "subject_cdr": subject_cdr,
        "query_aligned": aln["query_aligned"],
        "subject_aligned": aln["subject_aligned"],
        "mode": aln["mode"],
        "format": entry.get("format", ""),
        "genetics": entry.get("genetics", ""),
        "target": entry.get("target", ""),
        "highest_trial": entry.get("highest_trial", ""),
    }


def search_cdr(query_cdr: str, database: list, chain_key: str, cdr_name: str,
               top_n: int = 5, mode: str = "auto",
               min_score: float = 0.0, min_identity: float = 0.0) -> dict:
    """Find the top-N homologous database CDRs for one query CDR.

    Returns a dict with the ranked ``top_matches``, the number of candidates
    scanned/filtered, and the background identity distribution that gives the
    top hit a statistical context.
    """
    aln_mode = core.cdr_mode(cdr_name, mode)
    hits = []
    all_identities = []
    n_candidates = 0
    n_filtered = 0

    for entry, subject_cdr, source in _iter_chain_cdrs(database, chain_key, cdr_name):
        n_candidates += 1
        aln = core.align_sequences(query_cdr, subject_cdr, aln_mode)
        all_identities.append(aln["identity"])
        if aln["identity"] < min_identity or aln["score"] < min_score:
            n_filtered += 1
            continue
        hits.append(_make_hit(entry, subject_cdr, source, aln))

    hits.sort(key=lambda h: (-h["identity"], -h["score_norm"], -h["score"],
                             h["therapeutic"], h["chain_source"]))

    top = hits[:top_n]
    background = 0.0
    percentile = None
    if all_identities:
        background = round(statistics.median(all_identities), 4)
        if top:
            top_id = top[0]["identity"]
            at_least = sum(1 for i in all_identities if i >= top_id)
            percentile = round(at_least / len(all_identities), 4)

    return {
        "query_cdr": query_cdr,
        "query_len": len(query_cdr),
        "mode": aln_mode,
        "n_candidates": n_candidates,
        "n_filtered": n_filtered,
        "background_median_identity": background,
        "top_match_percentile": percentile,
        "top_match": top[0] if top else None,
        "top_matches": top,
    }


def run_query(queries: list, database: list, top_n: int = 5, mode: str = "auto",
              min_score: float = 0.0, min_identity: float = 0.0) -> list:
    """Run the CDR homology search for every query; collect per-chain errors."""
    all_results = []
    for name, vh_seq, vl_seq, meta in queries:
        print(f"  Processing {name} ...")
        kind, kind_info = core.chain_kind(vh_seq, vl_seq)
        result = {
            "name": name,
            "kind": kind,
            "kind_info": kind_info,
            "vh_seq": vh_seq,
            "vl_seq": vl_seq,
            "chain_source": meta.get("chain_source", {}),
            "cdrs": {},
            "matches": {},
            "errors": [],
        }
        for chain_key, seq in (("VH", vh_seq), ("VL", vl_seq)):
            if not seq:
                continue
            chain_type = chain_key[1]  # 'H' or 'L'
            posmap = core.number_sequence_imgt(seq, chain_type)
            if not posmap:
                result["errors"].append(
                    f"{chain_key}: ANARCI could not number this chain "
                    f"({len(seq)} aa); its CDRs are N/A")
                print(f"    [ERROR] {chain_key} could not be numbered by ANARCI",
                      file=sys.stderr)
                continue
            cdrs = core.extract_cdrs_from_posmap(posmap, chain_type)
            result["cdrs"][chain_key] = cdrs
            for cdr_name, cdr_seq in cdrs.items():
                if not cdr_seq:
                    result["errors"].append(
                        f"{chain_key}_{cdr_name}: empty after IMGT extraction")
                    continue
                full_key = f"{chain_key}_{cdr_name}"
                result["matches"][full_key] = search_cdr(
                    cdr_seq, database, chain_key, cdr_name,
                    top_n=top_n, mode=mode,
                    min_score=min_score, min_identity=min_identity)
        all_results.append(result)
    return all_results


# ---------------------------------------------------------------------------
# Statistics helpers
# ---------------------------------------------------------------------------

def _pct(value) -> str:
    return "N/A" if value is None else f"{value:.1%}"


def _cdr3_mean(result: dict):
    ids = [result["matches"][k]["top_match"]["identity"]
           for k in ("VH_CDR3", "VL_CDR3")
           if result["matches"].get(k, {}).get("top_match")]
    return sum(ids) / len(ids) if ids else None


def _all_cdr_mean(result: dict):
    ids = [m["top_match"]["identity"] for m in result["matches"].values()
           if m.get("top_match")]
    return sum(ids) / len(ids) if ids else None


def _truncate(text: str, width: int) -> str:
    text = str(text)
    return text if len(text) <= width else text[: width - 2] + ".."


def _top_drug_counts(database: list) -> list:
    """Count database entries per therapeutic name (fills appendix G)."""
    counts: dict = {}
    for entry in database:
        name = entry.get("name", "")
        if not name:
            continue
        counts[name] = counts.get(name, 0) + 1
    return sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))


def build_metadata(args, database: list, stats: dict) -> dict:
    """Provenance block: what was run, on which data, with which parameters."""
    meta = {
        "tool": core.TOOL_NAME,
        "tool_version": core.TOOL_VERSION,
        "generated_at": core.timestamp(),
        "python_version": core.python_version(),
        "libraries": core.library_versions(),
        "database": {
            "path": os.path.abspath(args.database),
            "sha256": core.file_sha256(args.database),
            "entries": len(database),
            "source_file": stats.get("source_file"),
            "source_sha256": stats.get("source_sha256"),
            "source_rows": stats.get("source_rows"),
            "preprocessed_at": stats.get("generated_at"),
        },
        "parameters": {
            "top": args.top,
            "align_mode": args.align,
            "min_score": args.min_score,
            "min_identity": args.min_identity,
            "cdr_boundaries_imgt": {k: list(v) for k, v in core.IMGT_CDR_RANGES.items()},
            "alignment": {
                "substitution_matrix": core.BLOSUM62,
                "open_gap_score": core.OPEN_GAP_SCORE,
                "extend_gap_score": core.EXTEND_GAP_SCORE,
                "mode_by_cdr": dict(core.DEFAULT_CDR_MODE),
            },
        },
    }
    return meta


# ---------------------------------------------------------------------------
# Markdown report
# ---------------------------------------------------------------------------

def _md_provenance(meta: dict) -> list:
    db = meta["database"]
    libs = ", ".join(f"{k} {v}" for k, v in meta["libraries"].items())
    return [
        "## 1. Provenance",
        "",
        "| Item | Value |",
        "|------|-------|",
        f"| Tool | {meta['tool']} v{meta['tool_version']} |",
        f"| Generated (UTC) | {meta['generated_at']} |",
        f"| Database | `{core.markdown_cell(db['path'])}` |",
        f"| Database SHA-256 | `{db['sha256'] or 'n/a'}` |",
        f"| Database entries | {db['entries']} |",
        f"| Source file | `{core.markdown_cell(db.get('source_file') or 'n/a')}` |",
        f"| Source SHA-256 | `{db.get('source_sha256') or 'n/a'}` |",
        f"| Source rows | {db.get('source_rows') if db.get('source_rows') is not None else 'n/a'} |",
        f"| Preprocessed (UTC) | {db.get('preprocessed_at') or 'n/a'} |",
        f"| Libraries | {libs} |",
        "",
    ]


def _md_method(meta: dict) -> list:
    p = meta["parameters"]
    return [
        "## 2. Method",
        "",
        "- **CDR definition**: IMGT (CDR1 27-38, CDR2 56-65, CDR3 105-117), "
        "numbered with ANARCI.",
        f"- **Alignment**: BLOSUM62, open gap {p['alignment']['open_gap_score']}, "
        f"extend gap {p['alignment']['extend_gap_score']}.",
        "- **CDR1/CDR2**: global (Needleman-Wunsch) identity over the full "
        "alignment. **CDR3**: local (Smith-Waterman) identity over the "
        "homologous core, because CDR3 lengths vary from 3 to >30 residues "
        "and a global denominator would divide by the length mismatch.",
        "- **Coverage**: fraction of the *query* CDR that participates in the "
        "alignment (1.0 for a global alignment).",
        "- **Score norm**: raw BLOSUM62 score / query self-alignment score.",
        f"- **Reported hits**: score >= {p['min_score']} and identity >= "
        f"{p['min_identity']}. Filtered-out candidates are counted per CDR.",
        "- **Background**: median identity of the query CDR against every "
        "database CDR; the top-hit percentile is its rank in that distribution.",
        "- **Caveat**: sequence homology is not functional homology. A high "
        "CDR3 identity indicates a shared loop, not a shared epitope or "
        "mechanism; use it as a regulatory/developability signal only.",
        "",
    ]


def _md_summary(all_results: list) -> list:
    lines = [
        "## 3. Summary",
        "",
        "| # | Antibody | Type | VH_CDR3 best | VH_CDR3 id | VL_CDR3 best | VL_CDR3 id | CDR3 mean id |",
        "|---|----------|------|--------------|-----------|--------------|-----------|--------------|",
    ]
    for i, r in enumerate(all_results, 1):
        vh3 = r["matches"].get("VH_CDR3", {}).get("top_match")
        vl3 = r["matches"].get("VL_CDR3", {}).get("top_match")
        lines.append(
            f"| {i} | {core.markdown_cell(_truncate(r['name'], 30))} | {r['kind']} "
            f"| {core.markdown_cell(_truncate(vh3['therapeutic'], 25)) if vh3 else 'N/A'} "
            f"| {_pct(vh3['identity']) if vh3 else 'N/A'} "
            f"| {core.markdown_cell(_truncate(vl3['therapeutic'], 25)) if vl3 else 'N/A'} "
            f"| {_pct(vl3['identity']) if vl3 else 'N/A'} "
            f"| {_pct(_cdr3_mean(r))} |")
    lines += [
        "",
        "> `CDR3 mean id` is the mean of the VH and VL CDR3 top-hit identities. "
        "CDR1/CDR2 are highly conserved across the clinical database, so their "
        "identities are near-saturated and are **not** folded into the headline "
        "number; see the per-CDR tables for their background medians.",
        "",
    ]
    return lines


def _md_query(result: dict, index: int, top_n: int) -> list:
    lines = ["---", f"## 4.{index} {core.markdown_cell(result['name'])}", ""]
    if result.get("kind") in ("vhh", "vhh_suspect"):
        hallmark = ", ".join(result.get("kind_info", {}).get("vhh_hallmark", []))
        if result["kind"] == "vhh":
            lines += [f"**Chain type**: VHH (camelid single domain; FR2 hallmark "
                      f"{hallmark}). CDRs are labelled VH_* by convention.", ""]
        else:
            lines += ["**Chain type**: single heavy chain without the camelid "
                      "FR2 hallmark (`vhh_suspect`); confirm the format.", ""]
    if result.get("errors"):
        lines.append("**Errors**")
        lines.append("")
        for err in result["errors"]:
            lines.append(f"- {core.markdown_cell(err)}")
        lines.append("")

    lines += [
        "### CDR Sequences (IMGT numbering)",
        "",
        "| Region | Query CDR | Len | Top Match | Identity | Coverage | Score | Background median |",
        "|--------|-----------|-----|-----------|----------|----------|-------|-------------------|",
    ]
    for full_key in ("VH_CDR1", "VH_CDR2", "VH_CDR3", "VL_CDR1", "VL_CDR2", "VL_CDR3"):
        m = result["matches"].get(full_key)
        if not m:
            lines.append(f"| {full_key} | - | - | - | - | - | - | - |")
            continue
        tm = m["top_match"]
        if tm:
            top_name = core.markdown_cell(_truncate(tm["therapeutic"], 30))
            top_id = _pct(tm["identity"])
            top_cov = _pct(tm["coverage"])
            top_score = f"{tm['score']:.0f}"
        else:
            top_name = top_id = top_cov = top_score = "N/A"
        lines.append(
            f"| {full_key} | `{core.markdown_cell(m['query_cdr'])}` | {m['query_len']} "
            f"| {top_name} | {top_id} | {top_cov} | {top_score} "
            f"| {_pct(m['background_median_identity'])} |")
    lines.append("")

    for full_key in ("VH_CDR1", "VH_CDR2", "VH_CDR3", "VL_CDR1", "VL_CDR2", "VL_CDR3"):
        m = result["matches"].get(full_key)
        if not m:
            continue
        lines.append(f"### {full_key} -- Top {top_n} Matches "
                     f"({m['mode']}, {m['n_candidates']} candidates, "
                     f"{m['n_filtered']} filtered)")
        lines.append("")
        if not m["top_matches"]:
            lines += ["No candidate passed the score/identity thresholds.", ""]
            continue
        lines.append("| Rank | Therapeutic | Chain | Identity | Coverage | Score | Score norm | Subject CDR | Alignment |")
        lines.append("|------|-------------|-------|----------|----------|-------|------------|-------------|-----------|")
        for j, hit in enumerate(m["top_matches"][:top_n], 1):
            aln_str = ""
            if hit["query_aligned"] and hit["subject_aligned"]:
                aln_str = (f"`{core.markdown_cell(hit['query_aligned'])}` / "
                           f"`{core.markdown_cell(hit['subject_aligned'])}`")
            lines.append(
                f"| {j} | {core.markdown_cell(_truncate(hit['therapeutic'], 30))} "
                f"| {hit['chain_source']} | {_pct(hit['identity'])} "
                f"| {_pct(hit['coverage'])} | {hit['score']:.0f} "
                f"| {hit['score_norm']:.2f} "
                f"| `{core.markdown_cell(hit['subject_cdr'])}` | {aln_str} |")
        lines.append("")
    return lines


def _md_appendix(database: list, stats: dict) -> list:
    lines = ["---", "# Appendix: Clinical Antibody Database Summary (TheraSAbDab)", ""]

    def simple_table(title, counts):
        out = [f"## {title}", "", "| Value | Count |", "|-------|-------|"]
        for value, count in sorted(counts.items(), key=lambda kv: (-kv[1], str(kv[0]))):
            out.append(f"| {core.markdown_cell(value)} | {count} |")
        out.append("")
        return out

    lines += simple_table("A. Antibody Format Distribution", stats.get("format_counts", {}))
    lines += simple_table("B. Genetics / Humanization Strategy", stats.get("genetics_counts", {}))
    lines += simple_table("C. Light Chain Type", stats.get("vd_lc_counts", {}))
    lines += simple_table("D. Highest Clinical Trial Phase", stats.get("trial_counts", {}))

    lines += ["## E. CDR Length Distribution (IMGT numbering)", "",
              "| CDR | Min | Max | Mean | Median | N |",
              "|-----|-----|-----|------|--------|---|"]
    for key in ("VH_CDR1", "VH_CDR2", "VH_CDR3", "VL_CDR1", "VL_CDR2", "VL_CDR3"):
        s = stats.get("cdr_length_summary", {}).get(key)
        if s:
            lines.append(f"| {key} | {s['min']} | {s['max']} | {s['mean']} "
                         f"| {s['median']} | {s['count']} |")
        else:
            lines.append(f"| {key} | - | - | - | - | 0 |")
    lines.append("")

    lines += ["## F. Top 30 Therapeutic Targets", "", "| Target | Count |", "|--------|-------|"]
    for target, count in sorted(stats.get("target_counts", {}).items(),
                                key=lambda kv: (-kv[1], str(kv[0])))[:30]:
        lines.append(f"| {core.markdown_cell(target)} | {count} |")
    lines.append("")

    lines += ["## G. Top 30 Drugs by Entry Count", "", "| Drug | Entries |", "|------|---------|"]
    for name, count in _top_drug_counts(database)[:30]:
        lines.append(f"| {core.markdown_cell(name)} | {count} |")
    lines.append("")
    return lines


def generate_report(all_results: list, database: list, stats: dict, metadata: dict,
                    output_dir: str, top_n: int) -> str:
    """Write the Markdown report; return its path."""
    os.makedirs(output_dir, exist_ok=True)
    lines = ["# CDR Homology Analysis Report", ""]
    lines.append(f"**Queries analysed:** {len(all_results)}  ")
    lines.append(f"**Database size:** {len(database)} clinical antibodies  ")
    n_errors = sum(len(r.get("errors", [])) for r in all_results)
    lines.append(f"**Warnings/errors:** {n_errors}")
    lines.append("")
    lines += _md_provenance(metadata)
    lines += _md_method(metadata)
    lines += _md_summary(all_results)
    for i, result in enumerate(all_results, 1):
        lines += _md_query(result, i, top_n)
    lines += _md_appendix(database, stats)

    path = os.path.join(output_dir, "cdr_homology_report.md")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"  Markdown report written: {path}")
    return path


# ---------------------------------------------------------------------------
# CSV detail files
# ---------------------------------------------------------------------------

def write_csv_details(all_results: list, output_dir: str) -> list:
    """Write one CSV per query; filenames are collision-free."""
    csv_dir = os.path.join(output_dir, "detail")
    os.makedirs(csv_dir, exist_ok=True)
    used: set = set()
    written = []
    for result in all_results:
        stem = core.safe_filename(result["name"], used)
        path = os.path.join(csv_dir, f"{stem}.csv")
        with open(path, "w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow([
                "CDR", "Query CDR", "Query Len", "Rank", "Therapeutic",
                "Chain Source", "Identity", "Identity (global)", "Coverage",
                "Score", "Score Norm", "Subject CDR", "Format", "Genetics",
                "Target", "Highest Trial", "Alignment",
            ])
            for full_key in ("VH_CDR1", "VH_CDR2", "VH_CDR3",
                             "VL_CDR1", "VL_CDR2", "VL_CDR3"):
                m = result["matches"].get(full_key)
                if not m:
                    continue
                for rank, hit in enumerate(m["top_matches"], 1):
                    writer.writerow([
                        full_key, m["query_cdr"], m["query_len"], rank,
                        hit["therapeutic"], hit["chain_source"],
                        f"{hit['identity']:.4f}", f"{hit['identity_global']:.4f}",
                        f"{hit['coverage']:.4f}", hit["score"], hit["score_norm"],
                        hit["subject_cdr"], hit.get("format", ""),
                        hit.get("genetics", ""), hit.get("target", ""),
                        hit.get("highest_trial", ""),
                        f"{hit['query_aligned']} / {hit['subject_aligned']}",
                    ])
        written.append(path)
        print(f"  CSV written: {path}")
    return written


# ---------------------------------------------------------------------------
# Word report
# ---------------------------------------------------------------------------

def generate_docx_report(all_results: list, database: list, stats: dict,
                         metadata: dict, output_dir: str, top_n: int) -> str:
    """Write the Word report; mirrors the Markdown sections including G."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    def shade(cell, color):
        cell._element.get_or_add_tcPr().append(
            cell._element.get_or_add_tcPr().makeelement(
                qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto",
                              qn("w:fill"): color}))

    def add_row(table, values, header=False):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                paragraph.style = doc.styles["Normal"]
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    if header:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if header:
                shade(cell, "2F5496")
        return row

    def make_table(headers, rows):
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            shade(table.rows[0].cells[i], "2F5496")
            for paragraph in table.rows[0].cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(8)
        for row in rows:
            add_row(table, row)
        doc.add_paragraph()
        return table

    doc = Document()
    title = doc.add_heading("CDR Homology Analysis Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Queries analysed: {len(all_results)}")
    doc.add_paragraph(f"Database size: {len(database)} clinical antibodies")
    n_errors = sum(len(r.get("errors", [])) for r in all_results)
    doc.add_paragraph(f"Warnings/errors: {n_errors}")

    # 1. Provenance
    db = metadata["database"]
    libs = ", ".join(f"{k} {v}" for k, v in metadata["libraries"].items())
    doc.add_heading("1. Provenance", level=1)
    make_table(["Item", "Value"], [
        ["Tool", f"{metadata['tool']} v{metadata['tool_version']}"],
        ["Generated (UTC)", metadata["generated_at"]],
        ["Database", db["path"]],
        ["Database SHA-256", db["sha256"] or "n/a"],
        ["Database entries", db["entries"]],
        ["Source file", db.get("source_file") or "n/a"],
        ["Source SHA-256", db.get("source_sha256") or "n/a"],
        ["Source rows", db.get("source_rows") if db.get("source_rows") is not None else "n/a"],
        ["Preprocessed (UTC)", db.get("preprocessed_at") or "n/a"],
        ["Libraries", libs],
    ])

    # 2. Method
    p = metadata["parameters"]
    doc.add_heading("2. Method", level=1)
    for bullet in (
        "CDR definition: IMGT (CDR1 27-38, CDR2 56-65, CDR3 105-117), numbered with ANARCI.",
        f"Alignment: BLOSUM62, open gap {p['alignment']['open_gap_score']}, "
        f"extend gap {p['alignment']['extend_gap_score']}.",
        "CDR1/CDR2: global (Needleman-Wunsch) identity. CDR3: local (Smith-Waterman) "
        "identity over the homologous core.",
        "Coverage: fraction of the query CDR participating in the alignment.",
        "Score norm: raw BLOSUM62 score / query self-alignment score.",
        f"Reported hits: score >= {p['min_score']} and identity >= {p['min_identity']}.",
        "Background: median identity against every database CDR.",
        "Caveat: sequence homology is not functional homology.",
    ):
        doc.add_paragraph(bullet, style="List Bullet")

    # 3. Summary
    doc.add_heading("3. Summary", level=1)
    rows = []
    for i, r in enumerate(all_results, 1):
        vh3 = r["matches"].get("VH_CDR3", {}).get("top_match")
        vl3 = r["matches"].get("VL_CDR3", {}).get("top_match")
        rows.append([
            i, _truncate(r["name"], 30), r["kind"],
            _truncate(vh3["therapeutic"], 25) if vh3 else "N/A",
            _pct(vh3["identity"]) if vh3 else "N/A",
            _truncate(vl3["therapeutic"], 25) if vl3 else "N/A",
            _pct(vl3["identity"]) if vl3 else "N/A",
            _pct(_cdr3_mean(r)),
        ])
    make_table(["#", "Antibody", "Type", "VH_CDR3 Best", "VH_CDR3 Id",
                "VL_CDR3 Best", "VL_CDR3 Id", "CDR3 Mean Id"], rows)
    doc.add_paragraph("CDR3 mean id aggregates only the two CDR3 top hits; "
                      "CDR1/CDR2 are near-saturated across the clinical database.")

    # 4. Per query
    for index, result in enumerate(all_results, 1):
        doc.add_heading(f"4.{index} {result['name']}", level=1)
        if result.get("kind") == "vhh":
            doc.add_paragraph("Chain type: VHH (camelid single domain).")
        elif result.get("kind") == "vhh_suspect":
            doc.add_paragraph("Chain type: single heavy chain without the camelid "
                              "FR2 hallmark (vhh_suspect); confirm the format.")
        if result.get("errors"):
            doc.add_heading("Errors", level=2)
            for err in result["errors"]:
                doc.add_paragraph(err, style="List Bullet")

        doc.add_heading("CDR Sequences (IMGT numbering)", level=2)
        rows = []
        for full_key in ("VH_CDR1", "VH_CDR2", "VH_CDR3",
                         "VL_CDR1", "VL_CDR2", "VL_CDR3"):
            m = result["matches"].get(full_key)
            if not m:
                rows.append([full_key, "-", "-", "-", "-", "-", "-", "-"])
                continue
            tm = m["top_match"]
            rows.append([
                full_key, m["query_cdr"], m["query_len"],
                _truncate(tm["therapeutic"], 30) if tm else "N/A",
                _pct(tm["identity"]) if tm else "N/A",
                _pct(tm["coverage"]) if tm else "N/A",
                f"{tm['score']:.0f}" if tm else "N/A",
                _pct(m["background_median_identity"]),
            ])
        make_table(["Region", "Query CDR", "Len", "Top Match", "Identity",
                    "Coverage", "Score", "Background median"], rows)

        for full_key in ("VH_CDR1", "VH_CDR2", "VH_CDR3",
                         "VL_CDR1", "VL_CDR2", "VL_CDR3"):
            m = result["matches"].get(full_key)
            if not m:
                continue
            doc.add_heading(
                f"{full_key} -- Top {top_n} Matches ({m['mode']}, "
                f"{m['n_candidates']} candidates, {m['n_filtered']} filtered)",
                level=2)
            if not m["top_matches"]:
                doc.add_paragraph("No candidate passed the score/identity thresholds.")
                continue
            rows = []
            for j, hit in enumerate(m["top_matches"][:top_n], 1):
                rows.append([
                    j, _truncate(hit["therapeutic"], 30), hit["chain_source"],
                    _pct(hit["identity"]), _pct(hit["coverage"]),
                    f"{hit['score']:.0f}", f"{hit['score_norm']:.2f}",
                    hit["subject_cdr"],
                    f"{hit['query_aligned']} / {hit['subject_aligned']}",
                ])
            make_table(["Rank", "Therapeutic", "Chain", "Identity", "Coverage",
                        "Score", "Score Norm", "Subject CDR", "Alignment"], rows)

    # Appendix
    doc.add_page_break()
    doc.add_heading("Appendix: Clinical Antibody Database Summary (TheraSAbDab)", level=1)

    def simple(title, counts):
        doc.add_heading(title, level=2)
        make_table(["Value", "Count"],
                   [[k, v] for k, v in sorted(counts.items(),
                                              key=lambda kv: (-kv[1], str(kv[0])))])

    simple("A. Antibody Format Distribution", stats.get("format_counts", {}))
    simple("B. Genetics / Humanization Strategy", stats.get("genetics_counts", {}))
    simple("C. Light Chain Type", stats.get("vd_lc_counts", {}))
    simple("D. Highest Clinical Trial Phase", stats.get("trial_counts", {}))

    doc.add_heading("E. CDR Length Distribution (IMGT numbering)", level=2)
    rows = []
    for key in ("VH_CDR1", "VH_CDR2", "VH_CDR3", "VL_CDR1", "VL_CDR2", "VL_CDR3"):
        s = stats.get("cdr_length_summary", {}).get(key)
        rows.append([key, s["min"], s["max"], s["mean"], s["median"], s["count"]]
                    if s else [key, "-", "-", "-", "-", 0])
    make_table(["CDR", "Min", "Max", "Mean", "Median", "N"], rows)

    doc.add_heading("F. Top 30 Therapeutic Targets", level=2)
    make_table(["Target", "Count"],
               [[t, c] for t, c in sorted(stats.get("target_counts", {}).items(),
                                          key=lambda kv: (-kv[1], str(kv[0])))[:30]])

    doc.add_heading("G. Top 30 Drugs by Entry Count", level=2)
    make_table(["Drug", "Entries"], [[n, c] for n, c in _top_drug_counts(database)[:30]])

    path = os.path.join(output_dir, "cdr_homology_report.docx")
    doc.save(path)
    print(f"  Word report written: {path}")
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _positive_int(text: str) -> int:
    try:
        value = int(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"invalid integer: {text!r}")
    if value < 1:
        raise argparse.ArgumentTypeError(f"must be >= 1 (got {value})")
    return value


def _fraction(text: str) -> float:
    try:
        value = float(text)
    except ValueError:
        raise argparse.ArgumentTypeError(f"invalid number: {text!r}")
    if not 0.0 <= value <= 1.0:
        raise argparse.ArgumentTypeError(f"must be within [0, 1] (got {value})")
    return value


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="CDR Homology Analysis: compare antibody CDRs against a "
                    "clinical antibody database")
    parser.add_argument("--database", default="data/therasabdab_cdrs.json",
                        help="Preprocessed CDR database JSON")
    parser.add_argument("--stats", default="data/therasabdab_stats.json",
                        help="Precomputed database statistics JSON")
    parser.add_argument("--input", "-i", help="FASTA file with query sequences")
    parser.add_argument("--vh", help="VH amino acid sequence")
    parser.add_argument("--vl", help="VL amino acid sequence")
    parser.add_argument("--name", help="Query name (single --vh/--vl mode)")
    parser.add_argument("--top", type=_positive_int, default=5,
                        help="Top N matches per CDR (default 5, must be >= 1)")
    parser.add_argument("--output", "-o", default="outputs/cdr_homology",
                        help="Output directory")
    parser.add_argument("--align", choices=("auto", "global", "local"), default="auto",
                        help="Alignment mode; 'auto' uses local for CDR3 and "
                             "global for CDR1/CDR2 (default: auto)")
    parser.add_argument("--min-score", type=float, default=0.0,
                        help="Minimum BLOSUM62 score for a hit to be reported "
                             "(default 0.0)")
    parser.add_argument("--min-identity", type=_fraction, default=0.0,
                        help="Minimum identity for a hit to be reported, 0-1 "
                             "(default 0.0)")
    parser.add_argument("--markdown", action="store_true", help="Write Markdown report")
    parser.add_argument("--json", action="store_true", help="Write raw JSON")
    parser.add_argument("--csv", action="store_true", help="Write per-query CSV detail")
    parser.add_argument("--docx", action="store_true", help="Write Word report")
    parser.add_argument("--all-formats", action="store_true",
                        help="Write every output format (default when no format "
                             "flag is given)")
    return parser


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if not args.input and not args.vh and not args.vl:
        parser.error("Provide --input FASTA or --vh/--vl sequences")

    if args.all_formats or not any(
            (args.markdown, args.json, args.csv, args.docx)):
        args.markdown = args.json = args.csv = args.docx = True

    if not os.path.exists(args.database):
        print(f"Error: database not found: {args.database}", file=sys.stderr)
        print("Run preprocess.py first:  python3 tools/cdr_homology/preprocess.py",
              file=sys.stderr)
        return 1

    print(f"Loading database from {args.database} ...")
    database = core.read_json(args.database)
    print(f"  {len(database)} clinical antibodies loaded")

    queries = parse_query_inputs(args)
    if not queries:
        print("Error: no queries found", file=sys.stderr)
        return 1
    print(f"  {len(queries)} query(ies) to process")

    print("Running CDR homology search ...")
    all_results = run_query(queries, database, top_n=args.top, mode=args.align,
                            min_score=args.min_score,
                            min_identity=args.min_identity)

    stats = core.read_json(args.stats) if os.path.exists(args.stats) else {}
    metadata = build_metadata(args, database, stats)

    os.makedirs(args.output, exist_ok=True)

    # JSON first: it is the machine-readable record and must survive a later
    # report-generation failure (e.g. python-docx missing).
    if args.json:
        path = core.write_json(
            os.path.join(args.output, "cdr_homology_results.json"),
            {"metadata": metadata, "queries": all_results})
        print(f"  Raw JSON written: {path}")

    if args.csv:
        write_csv_details(all_results, args.output)
    if args.markdown:
        generate_report(all_results, database, stats, metadata, args.output, args.top)
    if args.docx:
        try:
            generate_docx_report(all_results, database, stats, metadata,
                                 args.output, args.top)
        except ImportError as exc:
            print(f"  [WARN] Word report skipped: python-docx unavailable ({exc})",
                  file=sys.stderr)

    failures = sum(len(r.get("errors", [])) for r in all_results)
    if failures:
        print(f"\nCompleted with {failures} warning(s): some chains/CDRs could "
              f"not be analysed. See the report.", file=sys.stderr)
        return 2
    print("\nDone!")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
