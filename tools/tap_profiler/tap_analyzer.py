#!/usr/bin/env python3
"""TAP (Therapeutic Antibody Profiler) analysis tool.

Computes the five published TAP developability metrics (Ltot, PSH, PPC, PNC,
SFvCSP) on antibody structures or sequences. The metric definitions follow
Raybould et al., PNAS 2019 and Commun Biol 2024 - see tap_core.py for the
exact formulas and the differences from the legacy implementation.

Usage:
    python3 tools/tap_profiler/tap_analyzer.py --check
    python3 tools/tap_profiler/tap_analyzer.py --pdb structure.pdb
    python3 tools/tap_profiler/tap_analyzer.py --pdb structure.pdb \
        --sasa psa --psa-binary /path/to/psa --thresholds 2024
    python3 tools/tap_profiler/tap_analyzer.py --vh <VH_SEQ> --vl <VL_SEQ>
    python3 tools/tap_profiler/tap_analyzer.py --pdb-dir af3/ --all-formats

Requirements:
    ANARCI     IMGT numbering        (pip install anarci)
    FreeSASA   surface accessibility (pip install freesasa)
    JOY psa    optional, exact TAP Ala-X-Ala surface reference (--psa-binary)
    python-docx optional, Word reports
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from pathlib import Path
from typing import List, Optional

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from tap_core import (  # noqa: E402
    METRIC_ORDER,
    THRESHOLD_SETS,
    TAPResult,
    analyze_sequences,
    analyze_structure,
    _find_psa,
)

_STRUCTURE_EXTS = (".pdb", ".ent", ".cif", ".mmcif")


# ---------------------------------------------------------------------------
# Availability check
# ---------------------------------------------------------------------------

def check_tools() -> None:
    print("\n" + "=" * 60)
    print("  TAP Analyzer - Tool Availability Check")
    print("=" * 60)

    try:
        from scripts.humanize.anarci_adapter import is_anarci_available
        anarci_ok = is_anarci_available()
    except Exception:
        anarci_ok = False
    print("\n  Numbering:")
    print(f"    ANARCI IMGT:  {'✓ Available' if anarci_ok else '✗ Not installed'}")

    try:
        import freesasa  # noqa: F401
        freesasa_ok = True
    except ImportError:
        freesasa_ok = False
    try:
        psa_path = _find_psa()
    except FileNotFoundError:
        psa_path = None
    print("\n  Surface calculation:")
    print(f"    JOY psa:      {'✓ ' + psa_path if psa_path else '✗ Not on PATH (exact TAP reference)'}")
    print(f"    FreeSASA:     {'✓ Available' if freesasa_ok else '✗ Not installed'}")

    try:
        import docx  # noqa: F401
        docx_ok = True
    except ImportError:
        docx_ok = False
    print("\n  Optional:")
    print(f"    python-docx:  {'✓ Available' if docx_ok else '✗ Not installed (Word reports unavailable)'}")

    print("\n  Recommendations:")
    if not anarci_ok:
        print("    → pip install anarci          (IMGT numbering, required)")
    if not freesasa_ok and not psa_path:
        print("    → pip install freesasa        (surface accessibility, required)")
    if anarci_ok and freesasa_ok and not psa_path:
        print("    → optional: install JOY psa for the exact Ala-X-Ala surface")
        print("      reference used by the official TAP (the FreeSASA fallback")
        print("      agrees on ~93 % of surface/buried decisions)")
    print("")


# ---------------------------------------------------------------------------
# Report rendering
# ---------------------------------------------------------------------------

def _band(spec: dict) -> str:
    parts = []
    for lo, hi in spec.get("green", ()):
        if hi >= 1.0e8:
            parts.append(">= {}".format(_num(lo)))
        elif lo <= -1.0e8:
            parts.append("<= {}".format(_num(hi)))
        else:
            parts.append("{}-{}".format(_num(lo), _num(hi)))
    return ", ".join(parts) or "-"


def _num(v: float) -> str:
    if float(v) == int(v) and abs(v) < 1e6:
        return str(int(v))
    return "{:.2f}".format(v)


def format_text_report(result: TAPResult, title: str = "TAP Analysis") -> str:
    d = result.details
    lines = []
    lines.append("=" * 80)
    lines.append("  " + title)
    lines.append("=" * 80)
    lines.append("")
    lines.append(f"  Name:        {result.name}")
    if "pdb_file" in d:
        lines.append(f"  Structure:   {d['pdb_file']}")
        lines.append(f"  Chains:      VH={d['vh_chain']}  VL={d['vl_chain']}")
    lines.append(f"  Numbering:   {d['numbering']}")
    if "sasa_engine" in d:
        lines.append(f"  Surface:     {d['sasa_engine']}")
        lines.append(f"  Scope:       {'Fv only' if d['scope'] == 'fv' else 'whole complex'}")
    lines.append(f"  Thresholds:  {d['thresholds']}"
                 + ("  (2024, Raybould Commun Biol 2024)" if d['thresholds'] == "2024" else ""))
    if "vh_sequence" in d:
        lines.append(f"  VH length:   {len(d['vh_sequence'])} residues")
        lines.append(f"  VL length:   {len(d['vl_sequence'])} residues")

    lines.append("")
    lines.append("-" * 80)
    lines.append("  TAP METRICS")
    lines.append("-" * 80)
    lines.append("")
    lines.append("  {:<14s} {:<12s} {:<16s} {}".format("Metric", "Value", "Flag", "Green region"))
    lines.append("  {:<14s} {:<12s} {:<16s} {}".format(
        "-" * 14, "-" * 12, "-" * 16, "-" * 24))
    for metric in METRIC_ORDER:
        value = result.metrics.get(metric)
        if value is None:
            shown = "n/a"
        elif metric == "Ltot":
            shown = str(int(value))
        else:
            shown = "{:.3f}".format(value)
        flag = result.flags.get(metric, "unknown")
        band = _band(THRESHOLD_SETS[d["thresholds"]][metric])
        lines.append("  {:<14s} {:<12s} {:<16s} {}".format(
            metric, shown, _format_flag(flag), band))

    if d.get("cdr_lengths"):
        lines.append("")
        lines.append("-" * 80)
        lines.append("  CDR LENGTHS (IMGT)")
        lines.append("-" * 80)
        lines.append("")
        cols = list(d["cdr_lengths"].keys())
        header = "  {:<8s}".format("CDR") + "".join(
            "{:>10s}".format(c) for c in cols)
        lines.append(header)
        lines.append("  {:<8s}".format("-" * 8) + "".join(
            "{:>10s}".format("-" * 8) for _ in cols))
        for cdr in ("CDR1", "CDR2", "CDR3"):
            row = "  {:<8s}".format(cdr)
            for c in cols:
                row += "{:>10d}".format(d["cdr_lengths"][c].get(cdr, 0))
            lines.append(row)
        row = "  {:<8s}".format("Total")
        for c in cols:
            row += "{:>10d}".format(sum(d["cdr_lengths"][c].values()))
        lines.append(row)

    if "n_surface" in d:
        lines.append("")
        lines.append("-" * 80)
        lines.append("  STRUCTURE CONTEXT")
        lines.append("-" * 80)
        lines.append("")
        lines.append(f"  Surface residues:       {d['n_surface']}")
        lines.append(f"  CDR residues:           {d['n_cdr_residues']}")
        lines.append(f"  CDR vicinity residues:  {d['n_cdr_vicinity']}")
        lines.append(f"  Salt bridges:           {d['n_salt_bridges']}")
        lines.append(f"  Net charge VH / VL:     {d['net_charge_vh']:+.1f} / {d['net_charge_vl']:+.1f}"
                     f"   (SFvCSP = product)")

    lines.append("")
    lines.append("-" * 80)
    lines.append("  RISK ASSESSMENT")
    lines.append("-" * 80)
    lines.append("")
    lines.append(f"  Overall: {result.risk_summary}")
    reds = [m for m, f in result.flags.items() if f == "red"]
    ambers = [m for m, f in result.flags.items() if f == "amber"]
    unknown = [m for m, f in result.flags.items() if f == "unknown"]
    if reds:
        lines.append(f"  Red flags:   {', '.join(reds)}")
    if ambers:
        lines.append(f"  Amber flags: {', '.join(ambers)}")
    if unknown:
        lines.append(f"  Not evaluated (no structure): {', '.join(unknown)}")
    if d.get("warnings"):
        lines.append("")
        for w in d["warnings"]:
            lines.append(f"  NOTE: {w}")
    lines.append("")
    return "\n".join(lines)


def _format_flag(flag: str) -> str:
    return {"green": "✓ GREEN", "amber": "⚠ AMBER",
            "red": "✗ RED"}.get(flag, "? UNKNOWN")


def save_json_report(result: TAPResult, output_path: str) -> None:
    payload = {
        "name": result.name,
        "metrics": {k: result.metrics.get(k) for k in METRIC_ORDER},
        "flags": {k: result.flags.get(k) for k in METRIC_ORDER},
        "risk_summary": result.risk_summary,
        "details": result.details,
        "method": {
            "framework": "TAP (Therapeutic Antibody Profiler)",
            "references": [
                "Raybould et al. PNAS 2019;116(10):4025-4030",
                "Raybould et al. Commun Biol 2024;7:62",
            ],
        },
    }
    with open(output_path, "w") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)


def save_markdown_report(result: TAPResult, output_path: str,
                         title: str = "TAP Analysis") -> None:
    d = result.details
    lines = [f"# {title}", ""]
    lines.append(f"- **Name**: {result.name}")
    if "pdb_file" in d:
        lines.append(f"- **Structure**: `{d['pdb_file']}`")
        lines.append(f"- **Chains**: VH={d['vh_chain']}, VL={d['vl_chain']}")
    lines.append(f"- **Numbering**: {d['numbering']}")
    if "sasa_engine" in d:
        lines.append(f"- **Surface**: {d['sasa_engine']}")
        lines.append(f"- **Scope**: {d['scope']}")
    lines.append(f"- **Thresholds**: {d['thresholds']}")
    lines.append("")
    lines.append("## TAP metrics")
    lines.append("")
    lines.append("| Metric | Value | Flag | Green region |")
    lines.append("|--------|-------|------|--------------|")
    for metric in METRIC_ORDER:
        value = result.metrics.get(metric)
        shown = "n/a" if value is None else (
            str(int(value)) if metric == "Ltot" else "{:.3f}".format(value))
        band = _band(THRESHOLD_SETS[d["thresholds"]][metric])
        lines.append(f"| {metric} | {shown} | {_format_flag(result.flags.get(metric,'unknown'))} | {band} |")
    lines.append("")
    if d.get("cdr_lengths"):
        lines.append("## CDR lengths (IMGT)")
        lines.append("")
        cols = list(d["cdr_lengths"].keys())
        lines.append("| CDR | " + " | ".join(cols) + " |")
        lines.append("|-----|" + "|".join(["---"] * len(cols)) + "|")
        for cdr in ("CDR1", "CDR2", "CDR3"):
            lines.append(f"| {cdr} | " + " | ".join(
                str(d["cdr_lengths"][c].get(cdr, 0)) for c in cols) + " |")
        lines.append("| **Total** | " + " | ".join(
            str(sum(d["cdr_lengths"][c].values())) for c in cols) + " |")
        lines.append("")
    lines.append("## Risk assessment")
    lines.append("")
    lines.append(f"**{result.risk_summary}**")
    reds = [m for m, f in result.flags.items() if f == "red"]
    ambers = [m for m, f in result.flags.items() if f == "amber"]
    if reds:
        lines.append(f"- Red flags: {', '.join(reds)}")
    if ambers:
        lines.append(f"- Amber flags: {', '.join(ambers)}")
    if d.get("warnings"):
        lines.append("")
        for w in d["warnings"]:
            lines.append(f"> NOTE: {w}")
    lines.append("")
    with open(output_path, "w") as fh:
        fh.write("\n".join(lines))


def save_word_report(result: TAPResult, output_path: str,
                     title: str = "TAP Analysis") -> None:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word reports "
                           "(pip install python-docx)") from exc

    d = result.details
    doc = Document()
    doc.add_heading(title, level=0)
    p = doc.add_paragraph()
    p.add_run("Name: ").bold = True
    p.add_run(str(result.name))
    if "pdb_file" in d:
        p = doc.add_paragraph()
        p.add_run("Structure: ").bold = True
        p.add_run(f"{d['pdb_file']} (VH={d['vh_chain']}, VL={d['vl_chain']})")
    p = doc.add_paragraph()
    p.add_run("Numbering: ").bold = True
    p.add_run(str(d["numbering"]))
    if "sasa_engine" in d:
        p = doc.add_paragraph()
        p.add_run("Surface: ").bold = True
        p.add_run(f"{d['sasa_engine']}, scope={d['scope']}")
    p = doc.add_paragraph()
    p.add_run("Thresholds: ").bold = True
    p.add_run(str(d["thresholds"]))

    doc.add_heading("TAP metrics", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(("Metric", "Value", "Flag", "Green region")):
        hdr[i].text = text
        hdr[i].paragraphs[0].runs[0].bold = True
    for metric in METRIC_ORDER:
        value = result.metrics.get(metric)
        shown = "n/a" if value is None else (
            str(int(value)) if metric == "Ltot" else "{:.3f}".format(value))
        cells = table.add_row().cells
        cells[0].text = metric
        cells[1].text = shown
        cells[2].text = result.flags.get(metric, "unknown")
        cells[3].text = _band(THRESHOLD_SETS[d["thresholds"]][metric])

    if d.get("cdr_lengths"):
        doc.add_heading("CDR lengths (IMGT)", level=1)
        cols = list(d["cdr_lengths"].keys())
        table = doc.add_table(rows=1, cols=1 + len(cols))
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "CDR"
        for i, c in enumerate(cols, start=1):
            hdr[i].text = c
        for cdr in ("CDR1", "CDR2", "CDR3", "Total"):
            cells = table.add_row().cells
            cells[0].text = cdr
            for i, c in enumerate(cols, start=1):
                if cdr == "Total":
                    cells[i].text = str(sum(d["cdr_lengths"][c].values()))
                else:
                    cells[i].text = str(d["cdr_lengths"][c].get(cdr, 0))

    doc.add_heading("Risk assessment", level=1)
    para = doc.add_paragraph()
    run = para.add_run(result.risk_summary)
    run.bold = True
    run.font.size = Pt(12)
    reds = [m for m, f in result.flags.items() if f == "red"]
    ambers = [m for m, f in result.flags.items() if f == "amber"]
    if reds:
        doc.add_paragraph(f"Red flags: {', '.join(reds)}", style="List Bullet")
    if ambers:
        doc.add_paragraph(f"Amber flags: {', '.join(ambers)}", style="List Bullet")
    for w in d.get("warnings", []):
        doc.add_paragraph(f"NOTE: {w}", style="List Bullet")

    doc.add_paragraph("")
    doc.add_paragraph(
        "Metric definitions: Raybould et al., PNAS 2019 (five computational "
        "developability guidelines); thresholds: Raybould et al., Commun Biol "
        "2024 (TAP update).")
    doc.save(output_path)


# ---------------------------------------------------------------------------
# Output naming (keeps the legacy file-name convention)
# ---------------------------------------------------------------------------

def _output_stem(result: TAPResult) -> str:
    pdb_file = result.details.get("pdb_file")
    if pdb_file:
        pdb_path = Path(pdb_file)
        parent_name = pdb_path.parent.parent.name.replace(
            "Structure Prediction (Boltz-2)_", "")
        return f"{parent_name}_{pdb_path.stem}" if parent_name else pdb_path.stem
    return result.name


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="TAP (Therapeutic Antibody Profiler) analysis - "
                    "the five published developability guidelines",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Check available tools
  python3 tools/tap_profiler/tap_analyzer.py --check

  # Analyse a structure (default: ANARCI IMGT, FreeSASA, Fv scope, 2024 thresholds)
  python3 tools/tap_profiler/tap_analyzer.py --pdb structure.pdb

  # Exact TAP parity with the official JOY psa surface reference
  python3 tools/tap_profiler/tap_analyzer.py --pdb structure.pdb --sasa psa \\
      --psa-binary /path/to/psa

  # Structure already IMGT-numbered (no ANARCI lookup needed)
  python3 tools/tap_profiler/tap_analyzer.py --pdb structure.pdb --numbering native

  # Sequence-only (Ltot only)
  python3 tools/tap_profiler/tap_analyzer.py --vh SEQ --vl SEQ

  # Batch
  python3 tools/tap_profiler/tap_analyzer.py --pdb-dir af3/ --all-formats
""")
    parser.add_argument("--check", action="store_true", help="Check available tools")
    parser.add_argument("--pdb", help="Path to a PDB/CIF structure")
    parser.add_argument("--pdb-dir", help="Directory of structures (recursive)")
    parser.add_argument("--vh", help="VH amino acid sequence")
    parser.add_argument("--vl", help="VL amino acid sequence")
    parser.add_argument("--name", default="", help="Name for the analysis")
    parser.add_argument("--output", "-o", default="outputs/tap_profiler",
                        help="Output directory (default: outputs/tap_profiler)")
    parser.add_argument("--json", action="store_true", help="Save JSON report")
    parser.add_argument("--markdown", action="store_true", help="Save Markdown report")
    parser.add_argument("--word", action="store_true", help="Save Word report")
    parser.add_argument("--all-formats", action="store_true",
                        help="Save all formats (text + JSON + Markdown + Word)")
    parser.add_argument("--sasa", choices=("auto", "psa", "freesasa"),
                        default="auto",
                        help="Surface engine: auto = JOY psa if present, else "
                             "FreeSASA (default: auto)")
    parser.add_argument("--psa-binary", default=None,
                        help="Path to the JOY psa executable (exact TAP "
                             "Ala-X-Ala surface reference)")
    parser.add_argument("--numbering", choices=("anarci", "native"),
                        default="anarci",
                        help="anarci = number the sequence with ANARCI "
                             "(default); native = the structure is already "
                             "IMGT-numbered")
    parser.add_argument("--scope", choices=("fv", "complex"), default="fv",
                        help="fv = analyse Fv chains only (official TAP "
                             "protocol, default); complex = keep antigen/"
                             "constant chains for surface burial")
    parser.add_argument("--thresholds", choices=sorted(THRESHOLD_SETS),
                        default="2024",
                        help="Flag thresholds: 2024 (published, default), "
                             "2019, exscientia, legacy")
    parser.add_argument("--vh-chain", default=None,
                        help="Explicit PDB chain id for VH (skip auto-detect)")
    parser.add_argument("--vl-chain", default=None,
                        help="Explicit PDB chain id for VL (skip auto-detect)")

    args = parser.parse_args()

    if args.check:
        check_tools()
        return
    if not args.pdb and not args.pdb_dir and not args.vh:
        parser.error("Either --check, --pdb, --pdb-dir, or --vh/--vl must be specified")
    if args.vh and not args.vl:
        parser.error("--vl is required when using --vh")
    if (args.vh_chain and not args.vl_chain) or (args.vl_chain and not args.vh_chain):
        parser.error("--vh-chain and --vl-chain must be given together")

    os.makedirs(args.output, exist_ok=True)
    results: List[TAPResult] = []

    if args.pdb:
        print(f"\nAnalysing structure: {args.pdb}")
        try:
            result = analyze_structure(
                args.pdb, args.name or Path(args.pdb).stem,
                vh_chain=args.vh_chain, vl_chain=args.vl_chain,
                sasa_engine=args.sasa, psa_binary=args.psa_binary,
                numbering=args.numbering, scope=args.scope,
                thresholds=args.thresholds)
        except Exception as exc:
            print(f"  ERROR: {exc}", file=sys.stderr)
            sys.exit(1)
        results.append(result)
        print(format_text_report(result, f"TAP Analysis: {result.name}"))

    elif args.pdb_dir:
        files = [p for p in Path(args.pdb_dir).rglob("*")
                 if p.suffix.lower() in _STRUCTURE_EXTS]
        if not files:
            print("No PDB/CIF files found.", file=sys.stderr)
            sys.exit(1)
        print(f"\nFound {len(files)} structure files")
        for pdb_path in files:
            print(f"\nAnalysing: {pdb_path}")
            try:
                result = analyze_structure(
                    str(pdb_path), pdb_path.stem,
                    vh_chain=args.vh_chain, vl_chain=args.vl_chain,
                    sasa_engine=args.sasa, psa_binary=args.psa_binary,
                    numbering=args.numbering, scope=args.scope,
                    thresholds=args.thresholds)
            except Exception as exc:
                print(f"  ERROR: {exc}", file=sys.stderr)
                continue
            results.append(result)
            print(format_text_report(result, f"TAP Analysis: {result.name}"))

    else:
        print("\nAnalysing sequences (Ltot only; structure metrics need --pdb)")
        try:
            result = analyze_sequences(args.vh, args.vl,
                                       args.name or "sequence",
                                       thresholds=args.thresholds)
        except Exception as exc:
            print(f"  ERROR: {exc}", file=sys.stderr)
            sys.exit(1)
        results.append(result)
        print(format_text_report(result, f"TAP Analysis: {result.name}"))

    for result in results:
        stem = _output_stem(result)
        if args.json or args.all_formats:
            path = os.path.join(args.output, f"{stem}_tap.json")
            save_json_report(result, path)
            print(f"  JSON report saved: {path}")
        if args.markdown or args.all_formats:
            path = os.path.join(args.output, f"{stem}_tap.md")
            save_markdown_report(result, path, f"TAP Analysis: {result.name}")
            print(f"  Markdown report saved: {path}")
        if args.word or args.all_formats:
            path = os.path.join(args.output, f"{stem}_tap.docx")
            try:
                save_word_report(result, path, f"TAP Analysis: {result.name}")
                print(f"  Word report saved: {path}")
            except RuntimeError as exc:
                print(f"  Word report skipped: {exc}", file=sys.stderr)

    if len(results) > 1 and (args.json or args.all_formats):
        combined = os.path.join(args.output, "combined_tap_results.json")
        with open(combined, "w") as fh:
            json.dump([{
                "name": r.name,
                "metrics": r.metrics,
                "flags": r.flags,
                "risk_summary": r.risk_summary,
                "details": r.details,
            } for r in results], fh, indent=2, ensure_ascii=False)
        print(f"\nCombined JSON report saved: {combined}")


if __name__ == "__main__":
    main()
