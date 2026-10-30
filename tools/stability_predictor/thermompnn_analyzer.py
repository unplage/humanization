#!/usr/bin/env python3
"""ThermoMPNN stability analyzer for antibody back-mutations.

Predicts the folding-stability effect (ΔΔG) of single point mutations with
ThermoMPNN (Kuhlman Lab, MIT license) and, most importantly, maps ThermoMPNN's
0-based sequence positions to true Kabat/IMGT positions.

Two modes are supported:

  * ``ssm`` (default) - site-saturation mutagenesis at every position, or at a
    user supplied set of Kabat positions (``--positions``).
  * ``backmutation`` - evaluate exactly the human<->donor substitutions listed
    in a humanization pipeline report (``--pipeline-json``) or a
    ``backmutations_*.csv`` file (``--backmutations-csv``).  This is the mode
    that should be used to decide which back-mutations to keep.

Position handling
-----------------
ThermoMPNN's ``analysis/custom_inference.py`` reports ``position`` as a
**0-based index into the parsed chain sequence**, not PDB/Kabat numbering.
The wrapper therefore builds the same gapped sequence ThermoMPNN builds, maps
each index to the corresponding residue, and numbers the ungapped sequence with
the pipeline's Kabat/IMGT engines.  ``--positions H37`` always means Kabat 37.

Usage
-----
    # Evaluate every T1/T2 back-mutation on a V0 structure (recommended)
    python3 tools/stability_predictor/stability_analyzer.py \
        --pdb V0.pdb --chain all \
        --pipeline-json outputs/step3/humanization_result.json \
        --thermompnn-dir ~/ThermoMPNN \
        --python ~/miniconda3/envs/thermoMPNN/bin/python \
        --output outputs/stability --all-formats

    # Generic SSM restricted to true Kabat positions
    python3 tools/stability_predictor/stability_analyzer.py \
        --pdb structure.pdb --chain H \
        --positions 37,44,45,47,48,49,67,69,71,73,76,78,80,82,87,94 \
        --thermompnn-dir ~/ThermoMPNN \
        --python "conda run -n thermoMPNN python" \
        --output outputs/stability --json
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shlex
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass, asdict, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

# Repo root for optional Kabat/IMGT region annotation.
_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)

_TOOLS_DIR = os.path.dirname(os.path.abspath(__file__))

AA_ALPHABET = "ACDEFGHIKLMNPQRSTVWY"
CDR_REGIONS = {"CDR1", "CDR2", "CDR3"}

THREE_TO_ONE = {
    "ALA": "A", "CYS": "C", "ASP": "D", "GLU": "E", "PHE": "F", "GLY": "G",
    "HIS": "H", "ILE": "I", "LYS": "K", "LEU": "L", "MET": "M", "ASN": "N",
    "PRO": "P", "GLN": "Q", "ARG": "R", "SER": "S", "THR": "T", "VAL": "V",
    "TRP": "W", "TYR": "Y", "SEC": "U", "PYL": "O", "ASX": "B", "GLX": "Z",
    "XLE": "J", "XAA": "X", "MSE": "M",
}


# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------

@dataclass
class MutationResult:
    """Result for a single mutation (or a single back-mutation candidate)."""
    pdb_file: str
    chain: str
    position_kabat: str
    wildtype: str
    mutation: str
    ddG: Optional[float]
    region_kabat: str = ""
    region_imgt: str = ""
    position_seq: int = -1          # ThermoMPNN 0-based sequence index
    position_pdb: int = 0           # PDB author residue number
    insertion_code: str = ""
    is_stabilizing: bool = False
    is_significant: bool = False
    error: Optional[str] = None
    # Back-mutation specific fields
    tier: Optional[str] = None
    donor_aa: Optional[str] = None
    human_aa: Optional[str] = None
    direction: Optional[str] = None      # human_to_donor | donor_to_human
    ddG_backmutation: Optional[float] = None
    verdict: Optional[str] = None
    recommendation: Optional[str] = None
    note: Optional[str] = None


@dataclass
class StabilityReport:
    """Aggregated stability report for one structure/chain."""
    pdb_file: str
    chain: str
    chain_type: str = "H"
    mode: str = "ssm"                    # ssm | backmutation
    context: str = ""
    mutations: List[MutationResult] = field(default_factory=list)
    summary: Dict = field(default_factory=dict)
    error: Optional[str] = None
    warnings: List[str] = field(default_factory=list)


@dataclass
class ChainPosition:
    """One position of the chain as seen by ThermoMPNN and by the numberer."""
    seq_pos: int
    resseq: int
    insertion_code: str
    aa: str
    ungapped_index: int
    kabat_pos: Optional[str] = None
    kabat_region: Optional[str] = None
    imgt_pos: Optional[str] = None
    imgt_region: Optional[str] = None


@dataclass
class BackMutationCandidate:
    """A back-mutation listed by the humanization pipeline."""
    position: str
    donor_aa: str
    human_aa: str
    tier: str = ""
    composite: Optional[float] = None
    features: List[str] = field(default_factory=list)
    rationale: List[str] = field(default_factory=list)


@dataclass
class _PendingBackMutation:
    candidate: BackMutationCandidate
    cp: ChainPosition
    target_aa: str
    direction: str
    sign: int          # +1 when evaluating human->donor, -1 for donor->human


# ---------------------------------------------------------------------------
# Structure parsing (PDB / mmCIF) and ThermoMPNN-compatible sequence building
# ---------------------------------------------------------------------------

def _read_pdb_residues(pdb_path: str, chain: str) -> List[Tuple[int, str, str, str]]:
    """Return [(resseq, insertion_code, one_letter_aa, resname), ...].

    Mirrors ThermoMPNN's ``alt_parse_PDB_biounits``: ATOM records plus HETATM
    MSE (treated as MET) are used; duplicate altlocs/models are collapsed.
    """
    residues: List[Tuple[int, str, str, str]] = []
    seen = set()
    with open(pdb_path, "r", errors="ignore") as fh:
        for line in fh:
            rec = line[:6].strip()
            if rec not in ("ATOM", "HETATM"):
                continue
            resname = line[17:20].strip().upper()
            if rec == "HETATM" and resname != "MSE":
                continue
            if line[21] != chain:
                continue
            if resname == "MSE":
                resname = "MET"
            aa = THREE_TO_ONE.get(resname, "X")
            try:
                resseq = int(line[22:26])
            except ValueError:
                continue
            icode = line[26].strip()
            key = (resseq, icode)
            if key in seen:
                continue
            seen.add(key)
            residues.append((resseq, icode, aa, resname))
    residues.sort(key=lambda r: (r[0], r[1]))
    return residues


def _list_pdb_chains(pdb_path: str) -> List[str]:
    chains: List[str] = []
    seen = set()
    with open(pdb_path, "r", errors="ignore") as fh:
        for line in fh:
            rec = line[:6].strip()
            if rec not in ("ATOM", "HETATM"):
                continue
            resname = line[17:20].strip().upper()
            if rec == "HETATM" and resname != "MSE":
                continue
            ch = line[21]
            if ch not in seen:
                seen.add(ch)
                chains.append(ch)
    return chains


def _build_gapped_sequence(residues: List[Tuple[int, str, str, str]]):
    """Rebuild ThermoMPNN's gapped chain sequence.

    ThermoMPNN iterates every integer residue number between the chain's min
    and max and emits ``-`` for missing numbers. Insertion codes are sorted
    within a number. The returned ``seq_pos`` is exactly ThermoMPNN's
    ``position`` column.
    """
    by_num: Dict[int, List[Tuple[int, str, str, str]]] = {}
    for r in residues:
        by_num.setdefault(r[0], []).append(r)

    gapped: List[str] = []
    pos_info: List[Optional[Tuple[int, str, str, str]]] = []
    seqpos_to_ungapped: List[Optional[int]] = []
    ungapped: List[str] = []

    if not by_num:
        return "", pos_info, seqpos_to_ungapped, ""

    for n in range(min(by_num), max(by_num) + 1):
        group = sorted(by_num.get(n, []), key=lambda r: r[1])
        if not group:
            gapped.append("-")
            pos_info.append(None)
            seqpos_to_ungapped.append(None)
            continue
        for r in group:
            gapped.append(r[2])
            pos_info.append(r)
            seqpos_to_ungapped.append(len(ungapped))
            ungapped.append(r[2])

    return "".join(gapped), pos_info, seqpos_to_ungapped, "".join(ungapped)


def _fr4_score(seq: str, chain_type: str) -> int:
    """Small heuristic used only when both heavy and light numbering succeed."""
    tail = seq[-15:].upper()
    if chain_type == "H":
        if re.search(r"W.G.G.{0,4}VTVSS", tail):
            return 3
        if tail.endswith("VTVSS"):
            return 1
        return 0
    if re.search(r"[FL]G.GTK[LV]", tail):
        return 3
    if re.search(r"TK[LV][EK]IK$", tail) or tail.endswith("LTVL"):
        return 2
    return 0


def _number_sequence(seq: str, chain_type: str = "auto"):
    """Number the ungapped chain sequence in Kabat space.

    Returns (chain_type, NumberedChain|None, warnings).
    """
    from scripts.humanize.numbering import number_heavy, number_light

    warnings: List[str] = []
    if chain_type in ("H", "L"):
        fn = number_heavy if chain_type == "H" else number_light
        try:
            return chain_type, fn(seq), warnings
        except Exception as exc:  # noqa: BLE001
            return chain_type, None, [f"{chain_type} numbering failed: {exc}"]

    heavy = light = None
    heavy_err = light_err = None
    try:
        heavy = number_heavy(seq)
    except Exception as exc:  # noqa: BLE001
        heavy_err = str(exc)
    try:
        light = number_light(seq)
    except Exception as exc:  # noqa: BLE001
        light_err = str(exc)

    if heavy is not None and light is None:
        return "H", heavy, warnings
    if light is not None and heavy is None:
        return "L", light, warnings
    if heavy is not None and light is not None:
        hs, ls = _fr4_score(seq, "H"), _fr4_score(seq, "L")
        if hs >= ls:
            warnings.append(f"auto chain type: H (FR4 score {hs} vs {ls})")
            return "H", heavy, warnings
        warnings.append(f"auto chain type: L (FR4 score {hs} vs {ls})")
        return "L", light, warnings

    warnings.append(f"numbering failed (heavy: {heavy_err}; light: {light_err})")
    return "H", None, warnings


def _resolve_chain_id(pdb_path: str, chain_arg: str, chain_type: str = "auto") -> str:
    """Resolve a chain ID, allowing H/L to mean chain *type* for A/B structures."""
    available = _list_pdb_chains(pdb_path)
    if chain_arg in available:
        return chain_arg

    want = chain_arg.strip().upper()
    if want not in ("H", "L") and chain_type in ("H", "L"):
        want = chain_type

    if want in ("H", "L"):
        for cid in available:
            try:
                residues = _read_pdb_residues(pdb_path, cid)
                _, _, _, ungapped = _build_gapped_sequence(residues)
                if not ungapped:
                    continue
                ctype, _, _ = _number_sequence(ungapped, "auto")
            except Exception:  # noqa: BLE001
                continue
            if ctype == want:
                return cid

    raise ValueError(
        f"chain {chain_arg!r} not found in {os.path.basename(pdb_path)}; "
        f"available chains: {available}"
    )


def build_position_table(pdb_path: str, chain_arg: str, chain_type: str = "auto"):
    """Map ThermoMPNN sequence positions to true Kabat/IMGT positions.

    Returns (chain_id, chain_type, table, warnings, ungapped_sequence).
    ``table[seq_pos]`` is a ChainPosition or None for a ThermoMPNN gap.
    """
    chain_id = _resolve_chain_id(pdb_path, chain_arg, chain_type)
    residues = _read_pdb_residues(pdb_path, chain_id)
    if not residues:
        raise ValueError(
            f"no protein residues found for chain {chain_id!r} in {pdb_path}"
        )

    gapped, pos_info, seqpos_to_ungapped, ungapped = _build_gapped_sequence(residues)
    if not ungapped:
        raise ValueError(f"empty sequence for chain {chain_id!r}")

    ctype, numbered, warnings = _number_sequence(ungapped, chain_type)
    if numbered is None:
        raise ValueError("; ".join(warnings))

    idx_to_res: Dict[int, Any] = {r.index: r for r in numbered.residues}

    imgt_idx: Dict[int, Any] = {}
    try:
        from scripts.humanize.anarci_adapter import number_with_anarci_imgt
        imgt = number_with_anarci_imgt(ungapped, ctype)
        if imgt is not None:
            imgt_idx = {r.index: r for r in imgt.residues}
    except Exception:  # noqa: BLE001 - IMGT is optional
        pass

    table: List[Optional[ChainPosition]] = []
    unmapped = 0
    for seq_pos, info in enumerate(pos_info):
        if info is None:
            table.append(None)
            continue
        ui = seqpos_to_ungapped[seq_pos]
        nr = idx_to_res.get(ui) if ui is not None else None
        ir = imgt_idx.get(ui) if ui is not None else None
        if nr is None:
            unmapped += 1
        table.append(ChainPosition(
            seq_pos=seq_pos,
            resseq=info[0],
            insertion_code=info[1],
            aa=info[2],
            ungapped_index=ui if ui is not None else -1,
            kabat_pos=nr.pos if nr is not None else None,
            kabat_region=nr.region if nr is not None else None,
            imgt_pos=ir.pos if ir is not None else None,
            imgt_region=ir.region if ir is not None else None,
        ))
    if unmapped:
        warnings.append(f"{unmapped} residue(s) could not be mapped to Kabat numbering")
    if len(gapped) and gapped.count("-"):
        warnings.append(f"structure contains {gapped.count('-')} missing residue number(s)")

    return chain_id, ctype, table, warnings, ungapped


def prepare_structure(path: str) -> Tuple[str, Optional[str]]:
    """Return a PDB path usable by ThermoMPNN; convert mmCIF when needed.

    Returns (pdb_path, temp_path_to_delete_or_None).
    """
    suffix = Path(path).suffix.lower()
    if suffix not in (".cif", ".mmcif"):
        return path, None

    try:
        from Bio.PDB.MMCIFParser import MMCIFParser
        from Bio.PDB.PDBIO import PDBIO
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError(
            "biopython is required to read mmCIF input; install it or use a PDB file"
        ) from exc

    structure = MMCIFParser(QUIET=True).get_structure("input", path)
    tmp = tempfile.NamedTemporaryFile(suffix=".pdb", delete=False)
    tmp.close()
    try:
        io = PDBIO()
        io.set_structure(structure)
        io.save(tmp.name)
    except Exception as exc:  # noqa: BLE001
        os.unlink(tmp.name)
        raise RuntimeError(f"failed to convert {path} to PDB: {exc}") from exc
    return tmp.name, tmp.name


# ---------------------------------------------------------------------------
# ThermoMPNN installation / inference
# ---------------------------------------------------------------------------

def _python_parts(python_bin: str) -> List[str]:
    """Split a python path or a command such as ``conda run -n env python``."""
    raw = os.path.expanduser(python_bin.strip())
    if os.path.exists(raw):
        return [raw]
    parts = shlex.split(raw)
    if parts:
        parts[0] = os.path.expanduser(parts[0])
    return parts


def find_checkpoint(thermompnn_dir: str) -> Optional[Path]:
    """Find a ThermoMPNN stability checkpoint.

    Deliberately ignores ``vanilla_model_weights/``: those are ProteinMPNN
    weights and must never be used for ΔΔG prediction.
    """
    models = Path(thermompnn_dir).expanduser().resolve() / "models"
    cands = sorted(models.glob("*.ckpt")) + sorted(models.glob("*.pt"))
    if not cands:
        return None
    for c in cands:
        if "thermompnn" in c.name.lower():
            return c
    return cands[0]


def check_thermompnn_installation(
    thermompnn_dir: str,
    python_bin: str,
    timeout: int = 60,
) -> Tuple[bool, Optional[str]]:
    """Verify repo, checkpoint, and Python environment."""
    tm_dir = Path(thermompnn_dir).expanduser().resolve()

    for rel in ("analysis/custom_inference.py", "local.yaml"):
        if not (tm_dir / rel).exists():
            return False, f"Missing {rel} in {tm_dir}"

    if find_checkpoint(str(tm_dir)) is None:
        return False, (
            f"No ThermoMPNN checkpoint in {tm_dir}/models/ "
            "(expected thermoMPNN_default.pt/.ckpt; vanilla_model_weights are not stability models)"
        )

    parts = _python_parts(python_bin)
    if not parts:
        return False, "empty --python value"
    try:
        result = subprocess.run(
            parts + ["-c", (
                "import torch, pytorch_lightning, omegaconf, Bio, pandas; "
                "print('OK')"
            )],
            capture_output=True, text=True, timeout=timeout,
        )
        if result.returncode != 0:
            detail = (result.stderr or result.stdout or "").strip()
            return False, f"Python env check failed: {detail[-800:]}"
    except subprocess.TimeoutExpired:
        return False, "Python env check timed out"
    except FileNotFoundError:
        return False, (
            f"Python binary/command not found: {python_bin!r} "
            f"(tried {parts[0]!r})"
        )

    return True, None


def _run_subprocess(cmd: List[str], timeout: int, verbose: bool):
    if verbose:
        print(f"[ThermoMPNN] Running: {' '.join(cmd)}", file=sys.stderr)
    try:
        return subprocess.run(cmd, capture_output=not verbose, text=True, timeout=timeout), None
    except subprocess.TimeoutExpired:
        return None, (
            f"ThermoMPNN inference timed out (>{timeout}s). "
            "Increase --timeout or use a GPU/ThermoMPNN environment."
        )
    except FileNotFoundError as exc:
        return None, f"Failed to start ThermoMPNN command: {exc}"
    except Exception as exc:  # noqa: BLE001
        return None, f"Failed to run ThermoMPNN: {exc}"


def run_thermompnn_inference(
    pdb_path: str,
    chain: str,
    thermompnn_dir: str,
    python_bin: str,
    output_dir: str,
    timeout: int = 1800,
    verbose: bool = False,
    targets: Optional[List[Tuple[str, int, str]]] = None,
) -> Tuple[Optional[Path], Optional[str]]:
    """Run ThermoMPNN inference and return the output CSV path.

    ``targets`` is an optional list of ``(wildtype_aa, seq_pos, mutant_aa)``.
    When supplied, only those mutations are evaluated (fast targeted mode);
    otherwise full site-saturation mutagenesis is run.
    """
    tm_dir = Path(thermompnn_dir).expanduser().resolve()
    pdb_obj = Path(pdb_path).resolve()
    out_dir = Path(output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    ckpt = find_checkpoint(str(tm_dir))
    if ckpt is None:
        return None, f"No ThermoMPNN checkpoint found in {tm_dir}/models/"

    expected = out_dir / f"ThermoMPNN_inference_{pdb_obj.stem}.csv"
    if expected.exists():
        expected.unlink()

    parts = _python_parts(python_bin)
    if not parts:
        return None, "empty --python value"

    if targets is not None:
        if not targets:
            return None, None
        targets_json = out_dir / "targeted_mutations.json"
        targets_json.write_text(json.dumps([list(t) for t in targets]), encoding="utf-8")
        runner = Path(_TOOLS_DIR) / "targeted_inference.py"
        if not runner.exists():
            return None, f"missing helper script: {runner}"
        cmd = parts + [
            str(runner),
            "--thermompnn-dir", str(tm_dir),
            "--pdb", str(pdb_obj),
            "--chain", chain,
            "--model-path", str(ckpt),
            "--out-dir", str(out_dir),
            "--targets", str(targets_json),
        ]
    else:
        script = tm_dir / "analysis" / "custom_inference.py"
        cmd = parts + [
            str(script),
            "--pdb", str(pdb_obj),
            "--chain", chain,
            "--model_path", str(ckpt),
            "--out_dir", str(out_dir),
        ]

    start = time.time()
    result, err = _run_subprocess(cmd, timeout, verbose)
    if err:
        return None, err
    assert result is not None

    if result.returncode != 0:
        detail = ((result.stderr or "") + "\n" + (result.stdout or "")).strip()
        return None, f"ThermoMPNN failed (exit {result.returncode}): {detail[-2000:]}"

    if expected.exists():
        return expected, None

    # Fallback: a fresh CSV written by this run.
    fresh = [
        p for p in out_dir.glob("*.csv")
        if p.stat().st_mtime >= start - 2
    ]
    if fresh:
        return max(fresh, key=lambda p: p.stat().st_mtime), None
    return None, f"No output CSV found in {out_dir}"


# ---------------------------------------------------------------------------
# CSV parsing / annotation
# ---------------------------------------------------------------------------

def parse_thermompnn_csv(
    csv_path: str,
    pdb_name: str,
    chain: str,
) -> List[MutationResult]:
    """Parse ThermoMPNN's inference CSV.

    ``position`` is kept as ``position_seq`` (0-based); call
    :func:`apply_annotations` to attach Kabat/IMGT labels. Self-mutations are
    dropped (they are not mutations and always have ΔΔG ≈ 0).
    """
    mutations: List[MutationResult] = []
    with open(csv_path, "r", newline="") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            try:
                seq_pos = int(float(row["position"]))
                ddg = float(row["ddG_pred"])
                wt = str(row.get("wildtype", "")).strip().upper()
                mut = str(row.get("mutation", "")).strip().upper()
            except (KeyError, TypeError, ValueError):
                continue
            if not wt or not mut or wt == mut or seq_pos < 0:
                continue
            mutations.append(MutationResult(
                pdb_file=pdb_name,
                chain=chain,
                position_kabat="",
                wildtype=wt,
                mutation=mut,
                ddG=ddg,
                position_seq=seq_pos,
                is_stabilizing=ddg < 0,
                is_significant=abs(ddg) >= 1.0,
            ))
    return mutations


def apply_annotations(
    mutations: Sequence[MutationResult],
    table: Sequence[Optional[ChainPosition]],
    chain_type: str,
) -> None:
    """Attach true Kabat/IMGT position and region to each mutation in place."""
    for m in mutations:
        cp = table[m.position_seq] if 0 <= m.position_seq < len(table) else None
        if cp is None or not cp.kabat_pos:
            m.position_kabat = f"{chain_type}{m.position_seq + 1}"
            m.region_kabat = "-"
            m.region_imgt = "-"
            m.note = (m.note + "; " if m.note else "") + \
                "ThermoMPNN position could not be mapped to Kabat numbering"
            continue
        m.position_pdb = cp.resseq
        m.insertion_code = cp.insertion_code
        m.position_kabat = cp.kabat_pos
        m.region_kabat = cp.kabat_region or "-"
        m.region_imgt = cp.imgt_region or "-"


def _kabat_index(table: Sequence[Optional[ChainPosition]]) -> Dict[str, ChainPosition]:
    return {cp.kabat_pos: cp for cp in table if cp is not None and cp.kabat_pos}


def normalize_positions(
    positions: Optional[Sequence[str]],
    chain_type: str,
) -> List[str]:
    """Prefix un-prefixed position tokens with the chain type (H/L)."""
    normalized: List[str] = []
    for pos in positions or []:
        pos = str(pos).strip().upper()
        if not pos:
            continue
        if re.match(r"^[HL]\d", pos):
            normalized.append(pos)
        else:
            normalized.append(chain_type + pos)
    return normalized


def resolve_kabat_positions(
    table: Sequence[Optional[ChainPosition]],
    positions: Optional[Sequence[str]],
    chain_type: str,
) -> Tuple[set, List[str]]:
    """Map requested true Kabat positions to ThermoMPNN sequence indices."""
    if not positions:
        return set(), []
    by_pos = _kabat_index(table)
    wanted = set()
    missing: List[str] = []
    for pos in normalize_positions(positions, chain_type):
        if pos[0] != chain_type:
            continue
        wanted.add(pos)
    indices = set()
    for pos in sorted(wanted, key=_position_sort_key):
        cp = by_pos.get(pos)
        if cp is None:
            missing.append(pos)
        else:
            indices.add(cp.seq_pos)
    return indices, missing


def _position_sort_key(pos: str):
    m = re.match(r"^([A-Za-z]+)(\d+)([A-Za-z]*)$", pos or "")
    if m:
        return (m.group(1).upper(), int(m.group(2)), m.group(3).upper())
    return (pos or "", 0, "")


def parse_positions(pos_str: str, chain_type: str = "H") -> List[str]:
    """Parse comma-separated *Kabat* positions (e.g. ``37,44,H45,L24``)."""
    positions: List[str] = []
    for p in (pos_str or "").split(","):
        p = p.strip().upper()
        if not p:
            continue
        if re.match(r"^[HL]\d", p):
            positions.append(p)
        else:
            positions.append(chain_type + p)
    return positions


# ---------------------------------------------------------------------------
# Back-mutation candidates from the humanization pipeline
# ---------------------------------------------------------------------------

def parse_tiers(tiers: Optional[str]) -> Optional[set]:
    if not tiers or tiers.strip().lower() in ("all", "*"):
        return None
    return {t.strip().upper() for t in tiers.split(",") if t.strip()}


def load_backmutation_candidates(
    pipeline_json: Optional[str] = None,
    backmutations_csv: Optional[str] = None,
    pipeline_chain: Optional[str] = None,
    tiers: Optional[set] = None,
) -> List[BackMutationCandidate]:
    """Load back-mutation candidates from a pipeline JSON or CSV table."""
    candidates: List[BackMutationCandidate] = []

    if pipeline_json:
        with open(pipeline_json, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
        records = payload.get("chains", [])
        if pipeline_chain:
            wanted = pipeline_chain.strip()
            matched = [
                r for r in records
                if r.get("name") == wanted or r.get("chain_type") == wanted
            ]
            if matched:
                records = matched
        for rec in records:
            for b in rec.get("backmutations", []) or []:
                candidates.append(BackMutationCandidate(
                    position=str(b.get("position", "")).strip().upper(),
                    donor_aa=str(b.get("donor_aa", "")).strip().upper(),
                    human_aa=str(b.get("human_aa", "")).strip().upper(),
                    tier=str(b.get("tier", "")).strip().upper(),
                    composite=b.get("composite"),
                    features=list(b.get("features") or []),
                    rationale=list(b.get("rationale") or []),
                ))
    elif backmutations_csv:
        with open(backmutations_csv, "r", newline="") as fh:
            for row in csv.DictReader(fh):
                candidates.append(BackMutationCandidate(
                    position=str(row.get("position", "")).strip().upper(),
                    donor_aa=str(row.get("donor_aa", "")).strip().upper(),
                    human_aa=str(row.get("human_aa", "")).strip().upper(),
                    tier=str(row.get("tier", "")).strip().upper(),
                    composite=_to_float(row.get("composite")),
                    features=[f for f in str(row.get("features", "")).split("+") if f],
                    rationale=[r for r in str(row.get("rationale", "")).split(";") if r],
                ))

    if tiers:
        candidates = [c for c in candidates if c.tier.upper() in tiers]

    deduped: Dict[str, BackMutationCandidate] = {}
    for c in candidates:
        if c.position:
            deduped[c.position] = c
    return list(deduped.values())


def _to_float(value) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _prepare_backmutation_targets(
    table: Sequence[Optional[ChainPosition]],
    candidates: Sequence[BackMutationCandidate],
    chain_type: str,
) -> Tuple[List[_PendingBackMutation], List[Tuple[BackMutationCandidate, str]]]:
    """Split candidates into evaluable targets and skipped entries with reasons."""
    by_pos = _kabat_index(table)
    pending: List[_PendingBackMutation] = []
    skipped: List[Tuple[BackMutationCandidate, str]] = []

    for cand in sorted(candidates, key=lambda c: _position_sort_key(c.position)):
        if not cand.position.upper().startswith(chain_type):
            continue
        if cand.human_aa == "-" or cand.donor_aa == "-":
            skipped.append((cand, "FR indel: no single-residue substitution to evaluate"))
            continue
        cp = by_pos.get(cand.position)
        if cp is None:
            skipped.append((
                cand,
                "position not found in structure (numbering/structure mismatch)",
            ))
            continue
        if (cp.kabat_region or "") in CDR_REGIONS:
            skipped.append((
                cand,
                f"position maps to {cp.kabat_region}; CDR back-mutations are not allowed",
            ))
            continue

        wt = cp.aa
        if wt == cand.human_aa:
            target = cand.donor_aa
            direction = "human_to_donor"
            sign = 1
        elif wt == cand.donor_aa:
            target = cand.human_aa
            direction = "donor_to_human"
            sign = -1
        else:
            skipped.append((
                cand,
                f"structure residue {wt} matches neither human {cand.human_aa} "
                f"nor donor {cand.donor_aa}",
            ))
            continue

        if target not in AA_ALPHABET:
            skipped.append((cand, f"target residue {target!r} is not a standard amino acid"))
            continue

        pending.append(_PendingBackMutation(
            candidate=cand, cp=cp, target_aa=target,
            direction=direction, sign=sign,
        ))

    return pending, skipped


def _verdict(ddg_bm: Optional[float]) -> str:
    if ddg_bm is None:
        return "no prediction"
    if ddg_bm <= -1.0:
        return "strongly supports back-mutation"
    if ddg_bm <= -0.2:
        return "supports back-mutation"
    if ddg_bm < 0.2:
        return "neutral"
    if ddg_bm < 1.0:
        return "cautions against back-mutation"
    return "argues against back-mutation"


def _recommendation(tier: Optional[str], ddg_bm: Optional[float]) -> str:
    tier = (tier or "").upper()
    if ddg_bm is None:
        return "manual review"
    if ddg_bm <= -0.2:
        return "keep" if tier in ("T1", "T2") else "keep (optional)"
    if ddg_bm < 0.2:
        return "keep; structural evidence decides"
    if ddg_bm < 1.0:
        if tier == "T1":
            return "keep but flag stability risk"
        return "consider dropping"
    if tier == "T1":
        return "keep only if structurally essential; consider scaffold change"
    return "drop unless experimentally required"


def _result_for_pending(
    pending: _PendingBackMutation,
    row: Optional[MutationResult],
    pdb_name: str,
    chain: str,
) -> MutationResult:
    cand = pending.candidate
    cp = pending.cp
    ddg = row.ddG if row is not None else None
    ddg_bm = (row.ddG * pending.sign) if row is not None and row.ddG is not None else None
    return MutationResult(
        pdb_file=pdb_name,
        chain=chain,
        position_kabat=cp.kabat_pos or cand.position,
        wildtype=cp.aa,
        mutation=pending.target_aa,
        ddG=ddg,
        region_kabat=cp.kabat_region or "-",
        region_imgt=cp.imgt_region or "-",
        position_seq=cp.seq_pos,
        position_pdb=cp.resseq,
        insertion_code=cp.insertion_code,
        is_stabilizing=(ddg is not None and ddg < 0),
        is_significant=(ddg is not None and abs(ddg) >= 1.0),
        tier=cand.tier or None,
        donor_aa=cand.donor_aa,
        human_aa=cand.human_aa,
        direction=pending.direction,
        ddG_backmutation=ddg_bm,
        verdict=_verdict(ddg_bm),
        recommendation=_recommendation(cand.tier, ddg_bm),
        note=None if row is not None else "ThermoMPNN did not return this mutation",
    )


def _result_for_skipped(
    cand: BackMutationCandidate,
    reason: str,
    pdb_name: str,
    chain: str,
) -> MutationResult:
    return MutationResult(
        pdb_file=pdb_name,
        chain=chain,
        position_kabat=cand.position,
        wildtype="",
        mutation="",
        ddG=None,
        tier=cand.tier or None,
        donor_aa=cand.donor_aa,
        human_aa=cand.human_aa,
        error=reason,
        note=reason,
    )


# ---------------------------------------------------------------------------
# Summaries
# ---------------------------------------------------------------------------

def _summary_ssm(mutations: Sequence[MutationResult]) -> Dict:
    ddg_values = [m.ddG for m in mutations if m.ddG is not None]
    n_stab = sum(1 for v in ddg_values if v < -1e-9)
    n_dest = sum(1 for v in ddg_values if v > 1e-9)
    n_neutral = sum(1 for v in ddg_values if abs(v) <= 1e-9)
    return {
        "total": len(mutations),
        "stabilizing": n_stab,
        "destabilizing": n_dest,
        "neutral": n_neutral,
        "significant": sum(1 for v in ddg_values if abs(v) >= 1.0),
        "mean_ddg": sum(ddg_values) / len(ddg_values) if ddg_values else None,
        "min_ddg": min(ddg_values) if ddg_values else None,
        "max_ddg": max(ddg_values) if ddg_values else None,
    }


def _summary_backmutation(mutations: Sequence[MutationResult]) -> Dict:
    evaluated = [m for m in mutations if m.ddG_backmutation is not None]
    values: List[float] = []
    for m in evaluated:
        if m.ddG_backmutation is not None:
            values.append(m.ddG_backmutation)
    supports = sum(1 for v in values if v <= -0.2)
    neutral = sum(1 for v in values if -0.2 < v < 0.2)
    cautions = sum(1 for v in values if 0.2 <= v < 1.0)
    opposes = sum(1 for v in values if v >= 1.0)
    return {
        "candidates": len(mutations),
        "evaluated": len(evaluated),
        "skipped": len(mutations) - len(evaluated),
        "supports": supports,
        "neutral": neutral,
        "cautions": cautions,
        "opposes": opposes,
        "mean_ddg_backmutation": sum(values) / len(values) if values else None,
        "min_ddg_backmutation": min(values) if values else None,
        "max_ddg_backmutation": max(values) if values else None,
    }


def _context_label(pending: Sequence[_PendingBackMutation]) -> str:
    directions = {p.direction for p in pending}
    if directions == {"human_to_donor"}:
        return "human-framework context (V0-like)"
    if directions == {"donor_to_human"}:
        return "back-mutated context (V1/V2/V3-like)"
    if directions:
        return "mixed context"
    return "-"


# ---------------------------------------------------------------------------
# Text / Markdown / JSON / Word output
# ---------------------------------------------------------------------------

def _fmt(value: Optional[float], nd: int = 2) -> str:
    return "-" if value is None else f"{value:.{nd}f}"


def format_text_report(report: StabilityReport) -> str:
    if report.mode == "backmutation":
        return _format_backmutation_text(report)
    return _format_ssm_text(report)


def _format_ssm_text(report: StabilityReport) -> str:
    lines: List[str] = []
    lines.append("=" * 96)
    lines.append(
        f"  ThermoMPNN Stability Analysis: {report.pdb_file} "
        f"(chain {report.chain} / {report.chain_type})"
    )
    lines.append("=" * 96)

    if report.error:
        lines.append(f"\n  ERROR: {report.error}")
        for w in report.warnings:
            lines.append(f"  WARNING: {w}")
        return "\n".join(lines)

    if not report.mutations:
        lines.append("\n  No mutations analyzed (check --positions filter or structure).")
        return "\n".join(lines)

    lines.append(
        f"\n  {'Kabat':<8} {'PDB':<6} {'WT':<4} {'Mut':<4} "
        f"{'dG':>8} {'Effect':<14} {'Signif.':<8} "
        f"{'KabatReg':<9} {'IMGTReg':<9}"
    )
    lines.append("  " + "-" * 92)

    for m in sorted(report.mutations, key=lambda x: _position_sort_key(x.position_kabat)):
        effect = "STABILIZING" if m.is_stabilizing else "DESTABILIZING"
        signif = "YES" if m.is_significant else "no"
        insertion = m.insertion_code
        pdb_label = f"{m.position_pdb}{insertion}" if m.position_pdb else "-"
        lines.append(
            f"  {m.position_kabat:<8} {pdb_label:<6} {m.wildtype:<4} {m.mutation:<4} "
            f"{_fmt(m.ddG):>8} {effect:<14} {signif:<8} "
            f"{m.region_kabat or '-':<9} {m.region_imgt or '-':<9}"
        )

    s = report.summary
    lines.append(f"\n{'='*96}")
    lines.append("  SUMMARY")
    lines.append(f"{'='*96}")
    lines.append(f"  Total mutations analyzed:  {s.get('total', 0)}")
    lines.append(f"  Stabilizing (dG < 0):      {s.get('stabilizing', 0)}")
    lines.append(f"  Destabilizing (dG > 0):    {s.get('destabilizing', 0)}")
    lines.append(f"  Neutral (dG ~ 0):          {s.get('neutral', 0)}")
    lines.append(f"  Significant (|dG| >= 1.0): {s.get('significant', 0)}")
    if s.get("mean_ddg") is not None:
        lines.append(f"  Mean dG:                   {s['mean_ddg']:.2f} kcal/mol")
        lines.append(f"  Min dG:                    {s['min_ddg']:.2f} kcal/mol")
        lines.append(f"  Max dG:                    {s['max_ddg']:.2f} kcal/mol")

    for w in report.warnings:
        lines.append(f"  WARNING: {w}")
    lines.append("\n  Thresholds: |dG| >= 1.0 kcal/mol = significant; dG < 0 = stabilizing")
    lines.append("  Model: ThermoMPNN (Kuhlman Lab, MIT license)")
    return "\n".join(lines)


def _format_backmutation_text(report: StabilityReport) -> str:
    lines: List[str] = []
    lines.append("=" * 120)
    lines.append(
        f"  Back-mutation Stability Analysis: {report.pdb_file} "
        f"(chain {report.chain} / {report.chain_type})"
    )
    lines.append(f"  Context: {report.context}")
    lines.append("=" * 120)

    if report.error:
        lines.append(f"\n  ERROR: {report.error}")
        for w in report.warnings:
            lines.append(f"  WARNING: {w}")
        return "\n".join(lines)

    evaluated = [m for m in report.mutations if m.ddG_backmutation is not None]
    skipped = [m for m in report.mutations if m.ddG_backmutation is None]

    if evaluated:
        lines.append(
            f"\n  {'Kabat':<8} {'Tier':<5} {'H>D':<6} {'WT':<3} {'Mut':<4} "
            f"{'dG':>8} {'dG_bm':>8} {'Region':<9} {'Verdict':<30} Recommendation"
        )
        lines.append("  " + "-" * 116)
        for m in sorted(evaluated, key=lambda x: _position_sort_key(x.position_kabat)):
            hd = f"{m.human_aa}>{m.donor_aa}"
            lines.append(
                f"  {m.position_kabat:<8} {(m.tier or '-'):<5} {hd:<6} "
                f"{m.wildtype:<3} {m.mutation:<4} {_fmt(m.ddG):>8} "
                f"{_fmt(m.ddG_backmutation):>8} {m.region_kabat or '-':<9} "
                f"{m.verdict or '-':<30} {m.recommendation or '-'}"
            )

    if skipped:
        lines.append(f"\n  NOT EVALUATED ({len(skipped)})")
        lines.append("  " + "-" * 116)
        for m in sorted(skipped, key=lambda x: _position_sort_key(x.position_kabat)):
            lines.append(f"  {m.position_kabat:<8} {m.tier or '-':<5} {m.error or m.note or '-'}")

    s = report.summary
    lines.append(f"\n{'='*120}")
    lines.append("  SUMMARY")
    lines.append(f"{'='*120}")
    lines.append(f"  Candidates:                {s.get('candidates', 0)}")
    lines.append(f"  Evaluated:                 {s.get('evaluated', 0)}")
    lines.append(f"  Supports (dG_bm <= -0.2):  {s.get('supports', 0)}")
    lines.append(f"  Neutral:                   {s.get('neutral', 0)}")
    lines.append(f"  Cautions (0.2..1.0):       {s.get('cautions', 0)}")
    lines.append(f"  Opposes (dG_bm >= 1.0):    {s.get('opposes', 0)}")
    if s.get("mean_ddg_backmutation") is not None:
        lines.append(
            f"  Mean dG_backmutation:      {s['mean_ddg_backmutation']:.2f} kcal/mol"
        )

    for w in report.warnings:
        lines.append(f"  WARNING: {w}")
    lines.append("")
    lines.append("  dG_bm = predicted folding-stability effect of APPLYING the back-mutation")
    lines.append("          (negative = stabilizing). On a V0 structure dG is measured")
    lines.append("          directly as human->donor; on a V2/V3 structure it is measured as")
    lines.append("          donor->human and the sign is inverted, which is an approximation.")
    lines.append("  ThermoMPNN models monomer folding stability only: it does NOT predict")
    lines.append("  affinity, expression, aggregation, or VH/VL-interface effects.")
    return "\n".join(lines)


def save_json_report(report: StabilityReport, output_path: str) -> None:
    data = {
        "pdb_file": report.pdb_file,
        "chain": report.chain,
        "chain_type": report.chain_type,
        "mode": report.mode,
        "context": report.context,
        "mutations": [asdict(m) for m in report.mutations],
        "summary": report.summary,
        "error": report.error,
        "warnings": report.warnings,
    }
    if report.mode == "backmutation":
        data["backmutations"] = data["mutations"]
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)


def save_markdown_report(report: StabilityReport, output_path: str) -> None:
    if report.mode == "backmutation":
        _save_backmutation_markdown(report, output_path)
        return

    lines: List[str] = []
    lines.append(f"# ThermoMPNN Stability Analysis: {report.pdb_file}\n")
    lines.append(f"**Chain:** {report.chain} ({report.chain_type})  \n")

    if report.error:
        lines.append(f"\n**ERROR:** {report.error}\n")
    elif not report.mutations:
        lines.append("\nNo mutations analyzed.\n")
    else:
        lines.append(
            "| Kabat | PDB | WT | Mut | dG (kcal/mol) | Effect | Significant "
            "| Kabat Region | IMGT Region |"
        )
        lines.append("|-------|-----|-----|-----|----------------|--------|-------------|--------------|-------------|")
        for m in sorted(report.mutations, key=lambda x: _position_sort_key(x.position_kabat)):
            effect = "**STABILIZING**" if m.is_stabilizing else "DESTABILIZING"
            signif = "yes" if m.is_significant else ""
            pdb_label = f"{m.position_pdb}{m.insertion_code}" if m.position_pdb else "-"
            lines.append(
                f"| {m.position_kabat} | {pdb_label} | {m.wildtype} | {m.mutation} | "
                f"{_fmt(m.ddG)} | {effect} | {signif} | {m.region_kabat or '-'} "
                f"| {m.region_imgt or '-'} |"
            )

    s = report.summary
    lines.append("\n## Summary\n")
    lines.append(f"- **Total mutations:** {s.get('total', 0)}")
    lines.append(f"- **Stabilizing (dG < 0):** {s.get('stabilizing', 0)}")
    lines.append(f"- **Destabilizing (dG > 0):** {s.get('destabilizing', 0)}")
    lines.append(f"- **Neutral (dG ~ 0):** {s.get('neutral', 0)}")
    lines.append(f"- **Significant (|dG| >= 1.0):** {s.get('significant', 0)}")
    if s.get("mean_ddg") is not None:
        lines.append(f"- **Mean dG:** {s['mean_ddg']:.2f} kcal/mol")
        lines.append(f"- **Min dG:** {s['min_ddg']:.2f} kcal/mol")
        lines.append(f"- **Max dG:** {s['max_ddg']:.2f} kcal/mol")
    for w in report.warnings:
        lines.append(f"- **Warning:** {w}")
    with open(output_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def _save_backmutation_markdown(report: StabilityReport, output_path: str) -> None:
    lines: List[str] = []
    lines.append(f"# Back-mutation Stability Analysis: {report.pdb_file}\n")
    lines.append(f"**Chain:** {report.chain} ({report.chain_type})  ")
    lines.append(f"**Context:** {report.context}\n")

    if report.error:
        lines.append(f"\n**ERROR:** {report.error}\n")
    else:
        evaluated = [m for m in report.mutations if m.ddG_backmutation is not None]
        skipped = [m for m in report.mutations if m.ddG_backmutation is None]
        if evaluated:
            lines.append(
                "| Kabat | Tier | Human>Donor | WT | Mut | dG raw | **dG back-mutation** "
                "| Region | Verdict | Recommendation |"
            )
            lines.append("|-------|------|-------------|-----|-----|--------|-------------------------|--------|---------|----------------|")
            for m in sorted(evaluated, key=lambda x: _position_sort_key(x.position_kabat)):
                hd = f"{m.human_aa}>{m.donor_aa}"
                lines.append(
                    f"| {m.position_kabat} | {m.tier or '-'} | {hd} | {m.wildtype} "
                    f"| {m.mutation} | {_fmt(m.ddG)} | **{_fmt(m.ddG_backmutation)}** "
                    f"| {m.region_kabat or '-'} | {m.verdict or '-'} | {m.recommendation or '-'} |"
                )
        if skipped:
            lines.append("\n## Not evaluated\n")
            lines.append("| Kabat | Tier | Reason |")
            lines.append("|-------|------|--------|")
            for m in sorted(skipped, key=lambda x: _position_sort_key(x.position_kabat)):
                lines.append(f"| {m.position_kabat} | {m.tier or '-'} | {m.error or m.note or '-'} |")

    s = report.summary
    lines.append("\n## Summary\n")
    lines.append(f"- **Candidates:** {s.get('candidates', 0)}")
    lines.append(f"- **Evaluated:** {s.get('evaluated', 0)}")
    lines.append(f"- **Supports (dG_bm <= -0.2):** {s.get('supports', 0)}")
    lines.append(f"- **Neutral:** {s.get('neutral', 0)}")
    lines.append(f"- **Cautions (0.2..1.0):** {s.get('cautions', 0)}")
    lines.append(f"- **Opposes (dG_bm >= 1.0):** {s.get('opposes', 0)}")
    if s.get("mean_ddg_backmutation") is not None:
        lines.append(f"- **Mean dG_backmutation:** {s['mean_ddg_backmutation']:.2f} kcal/mol")
    for w in report.warnings:
        lines.append(f"- **Warning:** {w}")
    lines.append("")
    lines.append("> dG_bm is the predicted folding-stability effect of applying the back-mutation "
                 "(negative = stabilizing). On V2/V3 structures the sign is inferred by "
                 "inverting the donor->human prediction; prefer a V0 structure for direct "
                 "human->donor measurements. ThermoMPNN does not predict affinity, "
                 "expression, aggregation, or VH/VL-interface effects.")
    with open(output_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def save_word_report(report: StabilityReport, output_path: str) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError:
        print("  WARNING: python-docx not installed. Skipping Word report.", file=sys.stderr)
        print("  Install with: pip install python-docx", file=sys.stderr)
        return

    doc = Document()
    title = "Back-mutation Stability Analysis" if report.mode == "backmutation" \
        else "ThermoMPNN Stability Analysis"
    heading = doc.add_heading(f"{title}: {report.pdb_file}", 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Chain: {report.chain} ({report.chain_type})", style="Intense Quote")
    if report.mode == "backmutation" and report.context:
        doc.add_paragraph(f"Context: {report.context}", style="Intense Quote")

    if report.error:
        doc.add_paragraph(f"ERROR: {report.error}", style="Intense Quote")
        doc.save(output_path)
        return
    if not report.mutations:
        doc.add_paragraph("No mutations analyzed.")
        doc.save(output_path)
        return

    if report.mode == "backmutation":
        evaluated = [m for m in report.mutations if m.ddG_backmutation is not None]
        skipped = [m for m in report.mutations if m.ddG_backmutation is None]
        headers = ["Kabat", "Tier", "Human>Donor", "WT", "Mut", "dG raw",
                   "dG(bm)", "Region", "Verdict", "Recommendation"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(8)
        for m in sorted(evaluated, key=lambda x: _position_sort_key(x.position_kabat)):
            row = table.add_row()
            values = [
                m.position_kabat, m.tier or "-", f"{m.human_aa}>{m.donor_aa}",
                m.wildtype, m.mutation, _fmt(m.ddG), _fmt(m.ddG_backmutation),
                m.region_kabat or "-", m.verdict or "-", m.recommendation or "-",
            ]
            for i, val in enumerate(values):
                row.cells[i].text = str(val)
                for p in row.cells[i].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8)
            if m.ddG_backmutation is not None:
                color = RGBColor(0, 128, 0) if m.ddG_backmutation <= -0.2 else (
                    RGBColor(200, 0, 0) if m.ddG_backmutation >= 0.2 else RGBColor(128, 128, 0)
                )
                for p in row.cells[6].paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = color
                        run.font.bold = True
        if skipped:
            doc.add_heading("Not evaluated", level=2)
            t2 = doc.add_table(rows=1, cols=3)
            t2.style = "Light Grid Accent 1"
            for i, h in enumerate(["Kabat", "Tier", "Reason"]):
                t2.rows[0].cells[i].text = h
            for m in sorted(skipped, key=lambda x: _position_sort_key(x.position_kabat)):
                r = t2.add_row()
                r.cells[0].text = m.position_kabat
                r.cells[1].text = m.tier or "-"
                r.cells[2].text = m.error or m.note or "-"
    else:
        headers = ["Kabat", "PDB", "WT", "Mut", "dG", "Effect", "Significant",
                   "Kabat Region", "IMGT Region"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
        for m in sorted(report.mutations, key=lambda x: _position_sort_key(x.position_kabat)):
            row = table.add_row()
            pdb_label = f"{m.position_pdb}{m.insertion_code}" if m.position_pdb else "-"
            values = [
                m.position_kabat, pdb_label, m.wildtype, m.mutation, _fmt(m.ddG),
                "STABILIZING" if m.is_stabilizing else "DESTABILIZING",
                "yes" if m.is_significant else "",
                m.region_kabat or "-", m.region_imgt or "-",
            ]
            for i, val in enumerate(values):
                row.cells[i].text = str(val)
                for p in row.cells[i].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
            if m.is_significant:
                color = RGBColor(0, 128, 0) if m.is_stabilizing else RGBColor(255, 0, 0)
                for p in row.cells[5].paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = color
                        run.font.bold = True

    doc.add_heading("Summary", level=1)
    s = report.summary
    if report.mode == "backmutation":
        summary_items = [
            ("Candidates", s.get("candidates", 0)),
            ("Evaluated", s.get("evaluated", 0)),
            ("Supports", s.get("supports", 0)),
            ("Neutral", s.get("neutral", 0)),
            ("Cautions", s.get("cautions", 0)),
            ("Opposes", s.get("opposes", 0)),
        ]
        if s.get("mean_ddg_backmutation") is not None:
            summary_items.append((
                "Mean dG_backmutation",
                f"{s['mean_ddg_backmutation']:.2f} kcal/mol",
            ))
    else:
        summary_items = [
            ("Total mutations", s.get("total", 0)),
            ("Stabilizing (dG < 0)", s.get("stabilizing", 0)),
            ("Destabilizing (dG > 0)", s.get("destabilizing", 0)),
            ("Neutral (dG ~ 0)", s.get("neutral", 0)),
            ("Significant (|dG| >= 1.0)", s.get("significant", 0)),
        ]
        if s.get("mean_ddg") is not None:
            summary_items.extend([
                ("Mean dG", f"{s['mean_ddg']:.2f} kcal/mol"),
                ("Min dG", f"{s['min_ddg']:.2f} kcal/mol"),
                ("Max dG", f"{s['max_ddg']:.2f} kcal/mol"),
            ])

    summary_table = doc.add_table(rows=len(summary_items), cols=2)
    summary_table.style = "Light Grid Accent 1"
    for i, (label, value) in enumerate(summary_items):
        summary_table.rows[i].cells[0].text = label
        summary_table.rows[i].cells[1].text = str(value)

    doc.add_paragraph()
    for w in report.warnings:
        doc.add_paragraph(f"WARNING: {w}", style="Intense Quote")
    doc.add_paragraph(
        "dG_bm = effect of applying the back-mutation (negative = stabilizing). "
        "ThermoMPNN predicts folding stability only, not affinity/expression/interface effects.",
        style="Intense Quote",
    )
    doc.save(output_path)


# ---------------------------------------------------------------------------
# Main analysis
# ---------------------------------------------------------------------------

def _load_structure_context(pdb_path: str, chain_arg: str, chain_type: str):
    prepared, cleanup = prepare_structure(pdb_path)
    try:
        chain_id, ctype, table, warnings, ungapped = build_position_table(
            prepared, chain_arg, chain_type
        )
    finally:
        if cleanup:
            try:
                os.unlink(cleanup)
            except OSError:
                pass
    return chain_id, ctype, table, warnings, ungapped


def analyze_structure(
    pdb_path: str,
    chain: str,
    thermompnn_dir: str,
    python_bin: str,
    output_dir: str,
    positions: Optional[List[str]] = None,
    chain_type: str = "auto",
    candidates: Optional[Sequence[BackMutationCandidate]] = None,
    timeout: int = 1800,
    verbose: bool = False,
) -> StabilityReport:
    """Analyze one structure/chain for stability effects."""
    report = StabilityReport(pdb_file=os.path.basename(pdb_path), chain=chain)

    ok, err = check_thermompnn_installation(thermompnn_dir, python_bin)
    if not ok:
        report.error = f"ThermoMPNN not properly installed: {err}"
        return report

    prepared, cleanup = prepare_structure(pdb_path)
    try:
        try:
            chain_id, ctype, table, warnings, ungapped = build_position_table(
                prepared, chain, chain_type
            )
        except Exception as exc:  # noqa: BLE001
            report.error = f"Failed to number/parse structure: {exc}"
            return report

        report.chain = chain_id
        report.chain_type = ctype
        report.warnings.extend(warnings)

        by_pos = _kabat_index(table)
        pending: List[_PendingBackMutation] = []
        skipped: List[Tuple[BackMutationCandidate, str]] = []
        targets: Optional[List[Tuple[str, int, str]]] = None
        missing_positions: List[str] = []

        if candidates is not None:
            report.mode = "backmutation"
            chain_candidates = [
                c for c in candidates
                if c.position.upper().startswith(ctype)
            ]
            pending, skipped = _prepare_backmutation_targets(table, chain_candidates, ctype)
            report.context = _context_label(pending)
            targets = [(p.cp.aa, p.cp.seq_pos, p.target_aa) for p in pending]
        elif positions:
            normalized_positions = normalize_positions(positions, ctype)
            indices, missing_positions = resolve_kabat_positions(
                table, normalized_positions, ctype
            )
            if missing_positions:
                report.warnings.append(
                    "positions not found in structure: " + ", ".join(missing_positions)
                )
            targets = [
                (cp.aa, cp.seq_pos, aa)
                for cp in table
                if cp is not None and cp.seq_pos in indices
                for aa in AA_ALPHABET
                if aa != cp.aa
            ]

        csv_path, err = run_thermompnn_inference(
            prepared, chain_id, thermompnn_dir, python_bin, output_dir,
            timeout=timeout, verbose=verbose, targets=targets,
        )
        if err:
            report.error = err
            return report
        if csv_path is None:
            # No targets to evaluate (e.g. no candidate matches this chain).
            rows: List[MutationResult] = []
        else:
            rows = parse_thermompnn_csv(str(csv_path), report.pdb_file, chain_id)
            apply_annotations(rows, table, ctype)

        if candidates is not None:
            by_mutation = {(m.position_seq, m.mutation): m for m in rows}
            results: List[MutationResult] = []
            for p in pending:
                row = by_mutation.get((p.cp.seq_pos, p.target_aa))
                results.append(_result_for_pending(p, row, report.pdb_file, chain_id))
            for cand, reason in skipped:
                results.append(_result_for_skipped(cand, reason, report.pdb_file, chain_id))
            report.mutations = sorted(
                results, key=lambda x: _position_sort_key(x.position_kabat)
            )
            report.summary = _summary_backmutation(report.mutations)
        else:
            if positions:
                indices, _ = resolve_kabat_positions(
                    table, normalize_positions(positions, ctype), ctype
                )
                rows = [m for m in rows if m.position_seq in indices]
            if missing_positions:
                for pos in missing_positions:
                    report.warnings.append(f"requested Kabat position {pos} not in structure")
            report.mutations = sorted(
                rows, key=lambda x: (_position_sort_key(x.position_kabat), x.mutation)
            )
            report.summary = _summary_ssm(report.mutations)

        return report
    finally:
        if cleanup:
            try:
                os.unlink(cleanup)
            except OSError:
                pass


def analyze_batch(
    pdb_dir: str,
    chain: str,
    thermompnn_dir: str,
    python_bin: str,
    output_dir: str,
    positions: Optional[List[str]] = None,
    chain_type: str = "auto",
    candidates: Optional[Sequence[BackMutationCandidate]] = None,
    timeout: int = 1800,
    verbose: bool = False,
) -> List[StabilityReport]:
    """Analyze every PDB/CIF file in a directory."""
    pdb_dir_path = Path(pdb_dir)
    if not pdb_dir_path.exists():
        return [StabilityReport(
            pdb_file=pdb_dir, chain=chain,
            error=f"Directory not found: {pdb_dir}",
        )]

    files = sorted(
        p for p in pdb_dir_path.rglob("*")
        if p.is_file() and p.suffix.lower() in (".pdb", ".cif", ".mmcif")
    )
    if not files:
        return [StabilityReport(
            pdb_file=str(pdb_dir), chain=chain,
            error="No PDB/CIF files found",
        )]

    reports: List[StabilityReport] = []
    for pdb_file in files:
        for chain_id in expand_chain_arg(str(pdb_file), chain):
            print(f"\n[ThermoMPNN] Analyzing: {pdb_file} (chain {chain_id})")
            struct_out = Path(output_dir) / "thermompnn_raw" / pdb_file.stem / chain_id
            reports.append(analyze_structure(
                str(pdb_file), chain_id, thermompnn_dir, python_bin,
                str(struct_out), positions, chain_type, candidates,
                timeout, verbose,
            ))
    return reports


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def expand_chain_arg(pdb_path: str, chain_arg: str) -> List[str]:
    if chain_arg.strip().lower() == "all":
        prepared, cleanup = prepare_structure(pdb_path)
        try:
            return _list_pdb_chains(prepared)
        finally:
            if cleanup:
                try:
                    os.unlink(cleanup)
                except OSError:
                    pass
    return [c.strip() for c in chain_arg.split(",") if c.strip()]


def _report_stem(report: StabilityReport, multi_chain: bool) -> str:
    stem = Path(report.pdb_file).stem
    return f"{stem}_{report.chain}" if multi_chain else stem


def _save_reports(report: StabilityReport, args, multi_chain: bool):
    if not args.output:
        return
    stem = _report_stem(report, multi_chain)
    if args.json or args.all_formats:
        save_json_report(report, os.path.join(args.output, f"{stem}_stability.json"))
    if args.markdown or args.all_formats:
        save_markdown_report(report, os.path.join(args.output, f"{stem}_stability.md"))
    if args.word or args.all_formats:
        save_word_report(report, os.path.join(args.output, f"{stem}_stability.docx"))


def _write_combined(reports: Sequence[StabilityReport], output_dir: str) -> None:
    combined = {
        "reports": [
            {
                "pdb_file": r.pdb_file,
                "chain": r.chain,
                "chain_type": r.chain_type,
                "mode": r.mode,
                "context": r.context,
                "mutations": [asdict(m) for m in r.mutations],
                "summary": r.summary,
                "error": r.error,
                "warnings": r.warnings,
            }
            for r in reports
        ]
    }
    path = os.path.join(output_dir, "combined_stability_results.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(combined, fh, indent=2, ensure_ascii=False)
    print(f"\nCombined JSON saved: {path}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="ThermoMPNN stability / back-mutation analyzer",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Evaluate pipeline T1/T2 back-mutations on a V0 structure (recommended)
  %(prog)s --pdb V0.pdb --chain all \\
      --pipeline-json outputs/step3/humanization_result.json \\
      --thermompnn-dir ~/ThermoMPNN \\
      --python ~/miniconda3/envs/thermoMPNN/bin/python \\
      --output outputs/stability --all-formats

  # Generic SSM at true Kabat positions
  %(prog)s --pdb structure.pdb --chain H --positions 37,44,45,47 \\
      --thermompnn-dir ~/ThermoMPNN --python "conda run -n thermoMPNN python"

  # Batch a variant directory
  %(prog)s --pdb-dir variants/ --chain H \\
      --thermompnn-dir ~/ThermoMPNN --python ~/miniconda3/envs/thermoMPNN/bin/python \\
      --output outputs/stability/ --json
""",
    )
    parser.add_argument("--pdb", help="single PDB/CIF file to analyze")
    parser.add_argument("--pdb-dir", help="directory of PDB/CIF files for batch analysis")
    parser.add_argument(
        "--chain", default="H",
        help="chain ID, chain type (H/L), comma list, or 'all' (default: H)",
    )
    parser.add_argument(
        "--chain-type", default="auto", choices=("auto", "H", "L"),
        help="force heavy/light numbering; default auto-detect",
    )
    parser.add_argument(
        "--positions",
        help="comma-separated TRUE Kabat positions (e.g. '37,44,45' or 'H37,H35A')",
    )
    parser.add_argument(
        "--pipeline-json",
        help="humanization_result.json from the pipeline (targeted back-mutation mode)",
    )
    parser.add_argument(
        "--pipeline-chain",
        help="restrict --pipeline-json to a chain name (e.g. AMG110.H) or type (H/L)",
    )
    parser.add_argument(
        "--backmutations-csv",
        help="backmutations_*.csv from the pipeline (targeted back-mutation mode)",
    )
    parser.add_argument(
        "--tiers", default="all",
        help="tiers to evaluate in targeted mode: 'T1,T2', 'T1,T2,T3', or 'all' (default: all)",
    )
    parser.add_argument("--thermompnn-dir", required=True, help="path to ThermoMPNN repo")
    parser.add_argument(
        "--python", required=True,
        help="Python for ThermoMPNN env, e.g. '~/miniconda3/envs/thermoMPNN/bin/python' "
             "or 'conda run -n thermoMPNN python'",
    )
    parser.add_argument("--output", help="output directory for reports")
    parser.add_argument("--json", action="store_true", help="save JSON report")
    parser.add_argument("--markdown", action="store_true", help="save Markdown report")
    parser.add_argument("--word", action="store_true", help="save Word report")
    parser.add_argument("--all-formats", action="store_true", help="save all formats")
    parser.add_argument(
        "--timeout", type=int, default=1800,
        help="ThermoMPNN timeout in seconds (default: 1800)",
    )
    parser.add_argument("--verbose", action="store_true", help="verbose ThermoMPNN output")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if not args.pdb and not args.pdb_dir:
        parser.error("Either --pdb or --pdb-dir must be specified")
    if args.pdb and args.pdb_dir:
        parser.error("Use either --pdb or --pdb-dir, not both")
    if args.pipeline_json and args.backmutations_csv:
        parser.error("Use either --pipeline-json or --backmutations-csv, not both")
    if (args.pipeline_json or args.backmutations_csv) and args.positions:
        parser.error(
            "Use --positions for SSM mode or --pipeline-json/--backmutations-csv "
            "for back-mutation mode, not both"
        )

    positions = (
        [p.strip() for p in args.positions.split(",") if p.strip()]
        if args.positions else None
    )
    # Positions without an explicit H/L prefix are interpreted per-chain.

    candidates = None
    if args.pipeline_json or args.backmutations_csv:
        candidates = load_backmutation_candidates(
            pipeline_json=args.pipeline_json,
            backmutations_csv=args.backmutations_csv,
            pipeline_chain=args.pipeline_chain,
            tiers=parse_tiers(args.tiers),
        )
        print(f"[ThermoMPNN] Loaded {len(candidates)} back-mutation candidates")

    if args.output:
        os.makedirs(args.output, exist_ok=True)

    reports: List[StabilityReport] = []
    if args.pdb:
        chain_ids = expand_chain_arg(args.pdb, args.chain)
        multi = len(chain_ids) > 1
        base_out = args.output or "outputs/stability"
        for chain_id in chain_ids:
            # Keep raw CSVs per chain so multi-chain runs do not overwrite them.
            if multi:
                struct_out = os.path.join(
                    base_out, "thermompnn_raw", Path(args.pdb).stem, chain_id
                )
            else:
                struct_out = base_out
            report = analyze_structure(
                args.pdb, chain_id, args.thermompnn_dir, args.python,
                struct_out, positions,
                args.chain_type, candidates, args.timeout, args.verbose,
            )
            reports.append(report)
            print(format_text_report(report))
            _save_reports(report, args, multi)
    else:
        base = args.output or "outputs/stability"
        # Determine whether any file has multiple chains for naming.
        max_chains = 1
        for p in sorted(Path(args.pdb_dir).rglob("*")):
            if p.is_file() and p.suffix.lower() in (".pdb", ".cif", ".mmcif"):
                try:
                    max_chains = max(max_chains, len(expand_chain_arg(str(p), args.chain)))
                except Exception:  # noqa: BLE001
                    pass
        reports = analyze_batch(
            args.pdb_dir, args.chain, args.thermompnn_dir, args.python,
            base, positions, args.chain_type, candidates, args.timeout, args.verbose,
        )
        for report in reports:
            print(format_text_report(report))
            _save_reports(report, args, max_chains > 1)

    if len(reports) > 1 and args.output and (args.json or args.all_formats):
        _write_combined(reports, args.output)

    return 0


if __name__ == "__main__":
    sys.exit(main())
